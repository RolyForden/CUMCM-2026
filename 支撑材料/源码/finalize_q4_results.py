"""从第四问已完成的 CSV 生成两份官方模板副本，不重新运行全年模型。

输入（默认位于 outputs/q4）：
- q4_2_dispatch.csv / q4_2_daily_summary.csv
- q4_3_dispatch.csv / q4_3_daily_summary.csv / q4_3_versions.csv

输出：result4-2.xlsm、result4-3.xlsm、q4_workbook_audit.json。
脚本保留官方模板的原工作表、顺序、隐藏状态与 VBA，并只补题面要求但模板缺失的
“充放电量”和“紧急购电量”工作表。
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import shutil
from typing import Any
from zipfile import ZipFile

import openpyxl
import pandas as pd

from core.wallclock_output import WallclockDaySummary, aggregate_wallclock_days


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "outputs/q4"
TEMPLATE_DIR = ROOT / "data/raw/official/附件5"
EVAL_START = date(2025, 2, 1)
EVAL_END = date(2025, 12, 31)
EXPECTED_DAYS = [
    (EVAL_START + timedelta(days=i)).isoformat()
    for i in range((EVAL_END - EVAL_START).days + 1)
]
TOL = 1e-6
CD_BLOCK_LABELS = (
    "0:00-4:00", "4:00-8:00", "8:00-12:00",
    "12:00-16:00", "16:00-20:00", "20:00-24:00",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _vba_sha256(path: Path) -> str | None:
    with ZipFile(path) as archive:
        try:
            payload = archive.read("xl/vbaProject.bin")
        except KeyError:
            return None
    return hashlib.sha256(payload).hexdigest()


def _expected_frame(path: Path, required: set[str]) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"缺少正式中间产物：{path}")
    frame = pd.read_csv(path)
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"{path.name} 缺列：{sorted(missing)}")
    frame["date"] = frame["date"].astype(str).str[:10]
    if sorted(frame["date"].unique()) != EXPECTED_DAYS:
        raise ValueError(f"{path.name} 未精确覆盖 2025-02-01 至 2025-12-31")
    counts = frame.groupby("date").size()
    if not counts.eq(144).all():
        bad = counts[counts.ne(144)].to_dict()
        raise ValueError(f"{path.name} 存在非144槽日期：{bad}")
    ordered = frame.sort_values(["date", "slot"]).reset_index(drop=True)
    if any(group["slot"].astype(int).tolist() != list(range(144)) for _, group in ordered.groupby("date")):
        raise ValueError(f"{path.name} 槽位必须逐日唯一且为0至143")
    return ordered


def _daily_frame(path: Path, required: set[str]) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"缺少正式逐日汇总：{path}")
    frame = pd.read_csv(path)
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"{path.name} 缺列：{sorted(missing)}")
    frame["date"] = frame["date"].astype(str).str[:10]
    if frame["date"].duplicated().any() or frame["date"].tolist() != EXPECTED_DAYS:
        raise ValueError(f"{path.name} 日期必须按顺序唯一覆盖334天")
    return frame.set_index("date")


def _table_sheets(workbook: openpyxl.Workbook) -> list[str]:
    return [
        name for name in workbook.sheetnames
        if workbook[name].max_row == 335 and workbook[name].max_column == 147
    ]


def _check_template_dates(sheet: Any) -> None:
    for row, day_iso in enumerate(EXPECTED_DAYS, start=2):
        value = sheet.cell(row, 1).value
        if value is None or str(value)[:10] != day_iso:
            raise ValueError(f"{sheet.title} 第{row}行日期错位：{value!r}，应为{day_iso}")


def _time_label(ts: datetime, plan_day: str) -> str:
    plan_date = date.fromisoformat(plan_day)
    if ts.date() == plan_date:
        return ts.strftime("%H:%M")
    if ts.date() == plan_date + timedelta(days=1):
        return ts.strftime("%H:%M") + "+1"
    raise ValueError(f"{ts} 不属于计划日 {plan_day} 或其次日")


def _merge_emergency_intervals(dispatch: pd.DataFrame) -> list[dict[str, Any]]:
    """独立按相邻时间戳合并紧急购电槽。"""
    rows: list[dict[str, Any]] = []
    for day_iso, group in dispatch.groupby("date", sort=True):
        run: list[Any] = []

        def close_run() -> None:
            if not run:
                return
            start = pd.Timestamp(run[0].interval_start).to_pydatetime()
            end = pd.Timestamp(run[-1].interval_end).to_pydatetime()
            rows.append({
                "date": day_iso,
                "interval": f"{_time_label(start, day_iso)}-{_time_label(end, day_iso)}",
                "quantity": float(sum(float(item.grid_emergency) for item in run)),
            })

        for record in group.sort_values("slot").itertuples(index=False):
            if float(record.grid_emergency) > TOL:
                if run and pd.Timestamp(run[-1].interval_end) != pd.Timestamp(record.interval_start):
                    close_run()
                    run = []
                run.append(record)
            elif run:
                close_run()
                run = []
        close_run()
    return rows


def _load_wallclock_summaries(
    output_dir: Path, branch: str, dispatch: pd.DataFrame
) -> dict[str, WallclockDaySummary]:
    boundary_path = output_dir / f"{branch}_wallclock_boundary.csv"
    if not boundary_path.exists():
        raise FileNotFoundError(
            f"缺少{boundary_path.name}：现有正式CSV无法可靠恢复2月1日00:00边界，"
            "拒绝猜测或沿用错位分块"
        )
    boundary = pd.read_csv(boundary_path)
    summaries = aggregate_wallclock_days(
        dispatch.to_dict("records"),
        [date.fromisoformat(day) for day in EXPECTED_DAYS],
        boundary_records=boundary.to_dict("records"),
    )
    return {summary.day.isoformat(): summary for summary in summaries}


def _add_support_sheets(
    workbook: openpyxl.Workbook,
    dispatch: pd.DataFrame,
    wallclock: dict[str, WallclockDaySummary],
) -> tuple[int, int]:
    for name in ("充放电量", "紧急购电量"):
        if name in workbook.sheetnames:
            raise ValueError(f"官方模板已存在{name}，拒绝覆盖")

    cd = workbook.create_sheet("充放电量")
    header = ["日期"]
    for label in CD_BLOCK_LABELS:
        header.extend((f"{label}充电量", f"{label}放电量"))
    header.extend(("0:00储电量", "24:00储电量"))
    for column, value in enumerate(header, start=1):
        cd.cell(1, column, value)
    for row, day_iso in enumerate(EXPECTED_DAYS, start=2):
        summary = wallclock[day_iso]
        cd.cell(row, 1, day_iso)
        for block, values in enumerate(summary.blocks):
            cd.cell(row, 2 + 2 * block, values.charge)
            cd.cell(row, 3 + 2 * block, values.discharge)
        cd.cell(row, 14, summary.soc_0000)
        cd.cell(row, 15, summary.soc_2400)

    emergency = _merge_emergency_intervals(dispatch)
    em = workbook.create_sheet("紧急购电量")
    for column, value in enumerate(("日期", "紧急购电时间段", "紧急购电量"), start=1):
        em.cell(1, column, value)
    previous_day: str | None = None
    for row, item in enumerate(emergency, start=2):
        if item["date"] != previous_day:
            em.cell(row, 1, item["date"])
        em.cell(row, 2, item["interval"])
        em.cell(row, 3, item["quantity"])
        previous_day = item["date"]
    return cd.max_row, len(emergency)


def _write_table(
    sheet: Any,
    dispatch: pd.DataFrame,
    quantity_column: str,
    daily: pd.DataFrame,
    cost_column: str,
) -> None:
    _check_template_dates(sheet)
    for row, (day_iso, group) in enumerate(dispatch.groupby("date", sort=True), start=2):
        group = group.sort_values("slot")
        values = group[quantity_column].astype(float).to_numpy()
        for slot, value in enumerate(values):
            sheet.cell(row, 2 + slot, float(value))
        sheet.cell(row, 146, float(values.sum()))
        sheet.cell(row, 147, float(daily.loc[day_iso, cost_column]))


def _template_snapshot(workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    return [
        {
            "name": name,
            "state": workbook[name].sheet_state,
            "rows": workbook[name].max_row,
            "columns": workbook[name].max_column,
        }
        for name in workbook.sheetnames
    ]


def _finalize_q4_2(output_dir: Path) -> dict[str, Any]:
    dispatch = _expected_frame(
        output_dir / "q4_2_dispatch.csv",
        {"date", "slot", "interval_start", "interval_end", "grid_contract", "grid_emergency",
         "charge_actual", "discharge_actual", "soc_start", "soc_end"},
    )
    daily = _daily_frame(
        output_dir / "q4_2_daily_summary.csv", {"date", "cost_normal", "cost_total"}
    )
    wallclock = _load_wallclock_summaries(output_dir, "q4_2", dispatch)
    template = TEMPLATE_DIR / "result4-2.xlsm"
    target = output_dir / "result4-2.xlsm"
    shutil.copy2(template, target)
    workbook = openpyxl.load_workbook(target, keep_vba=True)
    before = _template_snapshot(workbook)
    tables = _table_sheets(workbook)
    visible = [name for name in tables if workbook[name].sheet_state == "visible"]
    if visible != ["计划购电量 (2)"]:
        workbook.close()
        raise ValueError(f"result4-2可见正式表异常：{visible}")
    _write_table(workbook[visible[0]], dispatch, "grid_contract", daily, "cost_normal")
    cd_rows, emergency_count = _add_support_sheets(workbook, dispatch, wallclock)
    workbook.save(target)
    workbook.close()
    return {
        "template": str(template.relative_to(ROOT)),
        "template_sha256": _sha256(template),
        "result_sha256": _sha256(target),
        "template_vba_sha256": _vba_sha256(template),
        "result_vba_sha256": _vba_sha256(target),
        "original_sheets": before,
        "quantity_sheet": visible[0],
        "quantity_semantics": "0:00制定的合同购电计划",
        "cost_semantics": "普通合同购电费；紧急购电在单独工作表列示",
        "charge_discharge_rows": cd_rows,
        "emergency_intervals": emergency_count,
    }


def _finalize_q4_3(output_dir: Path) -> dict[str, Any]:
    dispatch = _expected_frame(
        output_dir / "q4_3_dispatch.csv",
        {"date", "slot", "interval_start", "interval_end", "initial_contract", "final_contract",
         "grid_emergency", "charge_actual", "discharge_actual", "soc_start", "soc_end"},
    )
    daily = _daily_frame(
        output_dir / "q4_3_daily_summary.csv",
        {"date", "initial_contract_cost_actual", "total_cost_actual"},
    )
    wallclock = _load_wallclock_summaries(output_dir, "q4_3", dispatch)
    versions = output_dir / "q4_3_versions.csv"
    if not versions.exists() or pd.read_csv(versions, nrows=1).empty:
        raise ValueError("缺少非空的q4_3_versions.csv，不能只凭最终计划生成第三问对应结果")
    template = TEMPLATE_DIR / "result4-3.xlsm"
    target = output_dir / "result4-3.xlsm"
    shutil.copy2(template, target)
    workbook = openpyxl.load_workbook(target, keep_vba=True)
    before = _template_snapshot(workbook)
    tables = _table_sheets(workbook)
    hidden = [name for name in tables if workbook[name].sheet_state == "hidden"]
    visible = [name for name in tables if workbook[name].sheet_state == "visible"]
    if hidden != ["计划购电量"] or visible != ["计划购电量 (3)"]:
        workbook.close()
        raise ValueError(f"result4-3原始计划/最终计划表异常：hidden={hidden}, visible={visible}")
    _write_table(workbook[hidden[0]], dispatch, "initial_contract", daily, "initial_contract_cost_actual")
    _write_table(workbook[visible[0]], dispatch, "final_contract", daily, "total_cost_actual")
    cd_rows, emergency_count = _add_support_sheets(workbook, dispatch, wallclock)
    workbook.save(target)
    workbook.close()
    return {
        "template": str(template.relative_to(ROOT)),
        "template_sha256": _sha256(template),
        "result_sha256": _sha256(target),
        "template_vba_sha256": _vba_sha256(template),
        "result_vba_sha256": _vba_sha256(target),
        "original_sheets": before,
        "initial_plan_sheet": hidden[0],
        "final_plan_sheet": visible[0],
        "final_plan_sheet_semantics": "各槽最后一次已生效的调整后合同量；保留官方原名",
        "initial_cost_semantics": "0:00初始计划按目标槽真实价格计费",
        "final_cost_semantics": "初始计划、逐次调整与紧急购电的实际总费用",
        "charge_discharge_rows": cd_rows,
        "emergency_intervals": emergency_count,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    audit = {
        "result4-2": _finalize_q4_2(args.output_dir),
        "result4-3": _finalize_q4_3(args.output_dir),
    }
    audit["all_vba_preserved"] = all(
        row["template_vba_sha256"] == row["result_vba_sha256"]
        for row in (audit["result4-2"], audit["result4-3"])
    )
    (args.output_dir / "q4_workbook_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0 if audit["all_vba_preserved"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
