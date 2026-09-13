"""独立核验第四问两条正式结果。

本脚本不导入第四问生产回放器、价格预测器或工作簿生成器。它直接读取附件4、
正式 CSV、版本日志和结果工作簿，独立复算物理、实际价格费用、因果时间戳及模板。
"""

from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "outputs/q4"
PRICE_BOOK = ROOT / "data/raw/official/附件4.xlsm"
TEMPLATE_DIR = ROOT / "data/raw/official/附件5"
EVAL_START = date(2025, 2, 1)
EVAL_END = date(2025, 12, 31)
EXPECTED_DAYS = [
    (EVAL_START + timedelta(days=i)).isoformat()
    for i in range((EVAL_END - EVAL_START).days + 1)
]
TOL = 1e-6


def _vba_sha256(path: Path) -> str | None:
    with ZipFile(path) as archive:
        try:
            payload = archive.read("xl/vbaProject.bin")
        except KeyError:
            return None
    return hashlib.sha256(payload).hexdigest()


def _to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _load_actual_prices() -> dict[tuple[str, int], float]:
    workbook = openpyxl.load_workbook(PRICE_BOOK, read_only=True, data_only=True, keep_vba=True)
    candidates = [
        workbook[name] for name in workbook.sheetnames
        if workbook[name].max_row == 366 and workbook[name].max_column == 145
    ]
    if len(candidates) != 1:
        workbook.close()
        raise ValueError(f"附件4价格宽表数量={len(candidates)}，应为1")
    sheet = candidates[0]
    result: dict[tuple[str, int], float] = {}
    # read_only工作簿必须顺序迭代；反复cell()会从头扫描并退化为极慢的重复读取。
    for values in sheet.iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True):
        day_iso = _to_date(values[0]).isoformat()
        for slot, value in enumerate(values[1:]):
            if value is None or not np.isfinite(float(value)) or float(value) < 0:
                workbook.close()
                raise ValueError(f"附件4 {day_iso} slot={slot} 电价非法：{value!r}")
            result[(day_iso, slot)] = float(value)
    workbook.close()
    if len(result) != 365 * 144:
        raise AssertionError("附件4价格映射不完整")
    return result


def _read_dispatch(path: Path, q3_style: bool) -> pd.DataFrame:
    required = {
        "date", "slot", "interval_start", "interval_end", "decision_time",
        "price_forecast", "price_actual", "price_source_observed_at",
        "grid_contract", "grid_delivered", "grid_unused", "grid_emergency",
        "load_actual", "pv_available", "pv_used", "pv_curtail",
        "charge_planned", "charge_actual", "discharge_planned", "discharge_actual",
        "soc_start", "soc_end", "normal_cost_actual", "emergency_cost_actual",
    }
    if q3_style:
        required |= {"initial_contract", "final_contract"}
    frame = pd.read_csv(
        path,
        parse_dates=["interval_start", "interval_end", "decision_time", "price_source_observed_at"],
    )
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"{path.name} 缺列：{sorted(missing)}")
    frame["date"] = frame.date.astype(str).str[:10]
    return frame.sort_values(["date", "slot"]).reset_index(drop=True)


def _base_dispatch_checks(
    frame: pd.DataFrame, prices: dict[tuple[str, int], float], q3_style: bool
) -> tuple[dict[str, bool], dict[str, float]]:
    checks: dict[str, bool] = {}
    metrics: dict[str, float] = {}
    checks["shape_334x144"] = len(frame) == 334 * 144 and frame.groupby("date").size().eq(144).all()
    checks["dates_exact"] = sorted(frame.date.unique()) == EXPECTED_DAYS
    checks["slots_exact"] = all(
        group.slot.astype(int).tolist() == list(range(144))
        for _, group in frame.groupby("date", sort=True)
    )

    expected_price = np.asarray([
        prices[(str(row.date), int(row.slot))] for row in frame.itertuples(index=False)
    ])
    price_error = np.abs(frame.price_actual.to_numpy(float) - expected_price)
    metrics["max_actual_price_error"] = float(price_error.max(initial=0.0))
    checks["actual_price_from_attachment4"] = metrics["max_actual_price_error"] < TOL
    checks["forecast_price_finite_nonnegative"] = bool(
        np.isfinite(frame.price_forecast.to_numpy(float)).all()
        and (frame.price_forecast.to_numpy(float) >= 0).all()
    )

    checks["decision_before_target"] = bool((frame.decision_time <= frame.interval_start).all())
    checks["price_source_causal"] = bool((frame.price_source_observed_at <= frame.decision_time).all())
    expected_start = pd.to_datetime(frame.date) + pd.Timedelta(minutes=10) + pd.to_timedelta(frame.slot * 10, unit="m")
    expected_end = expected_start + pd.Timedelta(minutes=10)
    start_error = np.abs((frame.interval_start - expected_start).dt.total_seconds().to_numpy(float))
    end_error = np.abs((frame.interval_end - expected_end).dt.total_seconds().to_numpy(float))
    metrics["max_interval_label_error_seconds"] = float(max(start_error.max(initial=0), end_error.max(initial=0)))
    checks["scheme_a_intervals"] = metrics["max_interval_label_error_seconds"] < TOL
    gaps = (frame.interval_start.iloc[1:].reset_index(drop=True) - frame.interval_end.iloc[:-1].reset_index(drop=True)).dt.total_seconds()
    metrics["max_interval_gap_seconds"] = float(gaps.abs().max())
    checks["interval_continuity"] = metrics["max_interval_gap_seconds"] < TOL

    balance = (
        frame.grid_delivered + frame.grid_emergency + frame.pv_used + frame.discharge_actual
        - frame.load_actual - frame.charge_actual
    )
    soc_step = frame.soc_start + 0.9 * frame.charge_actual - frame.discharge_actual / 0.9 - frame.soc_end
    contract = frame.grid_contract - frame.grid_delivered - frame.grid_unused
    pv = frame.pv_available - frame.pv_used - frame.pv_curtail
    metrics["max_energy_residual"] = float(balance.abs().max())
    metrics["max_soc_residual"] = float(soc_step.abs().max())
    metrics["max_contract_residual"] = float(contract.abs().max())
    metrics["max_pv_residual"] = float(pv.abs().max())
    checks["energy_balance"] = metrics["max_energy_residual"] < 1e-8
    checks["soc_transition"] = metrics["max_soc_residual"] < 1e-8
    checks["contract_identity"] = metrics["max_contract_residual"] < TOL
    checks["pv_identity"] = metrics["max_pv_residual"] < TOL
    checks["all_flows_nonnegative"] = bool((frame[[
        "grid_contract", "grid_delivered", "grid_unused", "grid_emergency", "pv_available",
        "pv_used", "pv_curtail", "charge_actual", "discharge_actual",
    ]] >= -TOL).all().all())
    checks["soc_bounds"] = bool(
        frame[["soc_start", "soc_end"]].min().min() >= 1200 - TOL
        and frame[["soc_start", "soc_end"]].max().max() <= 10800 + TOL
    )
    checks["power_bounds"] = bool(
        frame[["charge_actual", "discharge_actual"]].max().max() <= 5000 / 6 + TOL
    )
    checks["no_simultaneous_charge_discharge"] = not bool(
        ((frame.charge_actual > TOL) & (frame.discharge_actual > TOL)).any()
    )
    checks["actual_not_above_planned"] = bool(
        (frame.charge_actual <= frame.charge_planned + TOL).all()
        and (frame.discharge_actual <= frame.discharge_planned + TOL).all()
    )
    starts = frame.groupby("date", sort=True).soc_start.first().to_numpy(float)
    ends = frame.groupby("date", sort=True).soc_end.last().to_numpy(float)
    metrics["max_cross_day_soc_error"] = float(np.abs(ends[:-1] - starts[1:]).max(initial=0.0))
    checks["cross_day_soc"] = metrics["max_cross_day_soc_error"] < TOL

    # 第三问对应流程的“正常费”只能指0:00初始计划费，不能对最终合同量再次收费。
    normal_quantity = frame.initial_contract if q3_style else frame.grid_contract
    normal = frame.price_actual * normal_quantity
    emergency = 5.0 * frame.price_actual * frame.grid_emergency
    metrics["max_normal_row_cost_error"] = float((normal - frame.normal_cost_actual).abs().max())
    metrics["max_emergency_row_cost_error"] = float((emergency - frame.emergency_cost_actual).abs().max())
    checks["normal_row_cost"] = metrics["max_normal_row_cost_error"] < TOL
    checks["emergency_row_cost"] = metrics["max_emergency_row_cost_error"] < TOL
    return checks, metrics


def _read_daily(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path)
    frame["date"] = frame.date.astype(str).str[:10]
    if frame.date.tolist() != EXPECTED_DAYS or frame.date.duplicated().any():
        raise ValueError(f"{path.name} 日期必须按顺序唯一覆盖334天")
    return frame.set_index("date")


def _check_q4_2_costs(dispatch: pd.DataFrame, daily: pd.DataFrame) -> tuple[dict[str, bool], dict[str, float]]:
    by_day = pd.DataFrame({
        "cost_normal": dispatch.price_actual * dispatch.grid_contract,
        "cost_emergency": 5.0 * dispatch.price_actual * dispatch.grid_emergency,
        "date": dispatch.date,
    }).groupby("date").sum()
    by_day["cost_total"] = by_day.cost_normal + by_day.cost_emergency
    checks: dict[str, bool] = {}
    metrics: dict[str, float] = {}
    for column in ("cost_normal", "cost_emergency", "cost_total"):
        if column not in daily.columns:
            raise ValueError(f"q4_2_daily_summary.csv 缺列：{column}")
        error = float((by_day[column] - daily[column]).abs().max())
        metrics[f"max_daily_{column}_error"] = error
        checks[f"daily_{column}"] = error < TOL
    return checks, metrics


def _check_q4_3_ledger(
    dispatch: pd.DataFrame, daily: pd.DataFrame, versions_path: Path
) -> tuple[dict[str, bool], dict[str, float], pd.DataFrame]:
    versions = pd.read_csv(
        versions_path,
        parse_dates=["issue_time", "price_source_observed_at"],
    )
    required = {
        "date", "slot", "version", "issue_time", "quantity", "price_forecast",
        "price_source_observed_at",
    }
    missing = required - set(versions.columns)
    if missing:
        raise ValueError(f"q4_3_versions.csv 缺列：{sorted(missing)}")
    versions["date"] = versions.date.astype(str).str[:10]
    dispatch_by_key = {
        (str(row.date), int(row.slot)): row for row in dispatch.itertuples(index=False)
    }
    version_ok = True
    final_ok = True
    causal_ok = True
    rows: list[dict[str, Any]] = []
    for key, group in versions.groupby(["date", "slot"], sort=False):
        if key not in dispatch_by_key:
            version_ok = False
            continue
        group = group.sort_values("version")
        record = dispatch_by_key[key]
        numbers = group.version.astype(int).tolist()
        version_ok &= numbers == list(range(len(group))) and group.issue_time.is_monotonic_increasing
        causal_ok &= bool((group.issue_time <= record.interval_start).all())
        causal_ok &= bool((group.price_source_observed_at <= group.issue_time).all())
        quantities = group.quantity.to_numpy(float)
        final_ok &= abs(float(quantities[0]) - float(record.initial_contract)) < TOL
        final_ok &= abs(float(quantities[-1]) - float(record.final_contract)) < TOL
        final_ok &= abs(float(record.final_contract) - float(record.grid_contract)) < TOL
        price = float(record.price_actual)
        delta = np.diff(quantities)
        initial = price * float(quantities[0])
        up = 1.5 * price * float(np.maximum(delta, 0).sum())
        down = 0.5 * price * float(np.maximum(-delta, 0).sum())
        emergency = 5.0 * price * float(record.grid_emergency)
        rows.append({
            "date": key[0], "initial_contract_cost_actual": initial,
            "up_cost_actual": up, "down_credit_actual": down,
            "adjustment_cashflow_actual": up - down,
            "emergency_cost_actual": emergency,
            "total_cost_actual": initial + up - down + emergency,
        })
    checks = {
        "version_keys_complete": len(rows) == len(dispatch) and len(versions.groupby(["date", "slot"])) == len(dispatch),
        "version_sequence": bool(version_ok),
        "version_causality": bool(causal_ok),
        "version_dispatch_match": bool(final_ok),
    }
    ledger = pd.DataFrame(rows).groupby("date").sum()
    metrics: dict[str, float] = {}
    mapping = (
        "initial_contract_cost_actual", "adjustment_cashflow_actual",
        "emergency_cost_actual", "total_cost_actual",
    )
    for column in mapping:
        if column not in daily.columns:
            raise ValueError(f"q4_3_daily_summary.csv 缺列：{column}")
        error = float((ledger[column] - daily[column]).abs().max())
        metrics[f"max_daily_{column}_error"] = error
        checks[f"daily_{column}"] = error < TOL
    checks["no_qfinal_double_charge"] = bool(
        np.allclose(
            ledger.total_cost_actual.to_numpy(float),
            (ledger.initial_contract_cost_actual + ledger.adjustment_cashflow_actual
             + ledger.emergency_cost_actual).to_numpy(float),
            rtol=0.0, atol=TOL,
        )
    )
    return checks, metrics, versions


def _time_label(ts: pd.Timestamp, plan_day: str) -> str:
    value = ts.strftime("%H:%M")
    if ts.date() == date.fromisoformat(plan_day) + timedelta(days=1):
        value += "+1"
    return value


def _expected_emergency(dispatch: pd.DataFrame) -> list[tuple[str, str, float]]:
    expected: list[tuple[str, str, float]] = []
    for day_iso, group in dispatch.groupby("date", sort=True):
        run: list[Any] = []

        def close_run() -> None:
            if not run:
                return
            expected.append((
                day_iso,
                f"{_time_label(run[0].interval_start, day_iso)}-{_time_label(run[-1].interval_end, day_iso)}",
                float(sum(float(item.grid_emergency) for item in run)),
            ))

        for row in group.sort_values("slot").itertuples(index=False):
            if float(row.grid_emergency) > TOL:
                if run and run[-1].interval_end != row.interval_start:
                    close_run()
                    run = []
                run.append(row)
            elif run:
                close_run()
                run = []
        close_run()
    return expected


def _sheet_snapshot(workbook: openpyxl.Workbook) -> list[tuple[str, str]]:
    return [(name, workbook[name].sheet_state) for name in workbook.sheetnames]


def _check_workbook(
    name: str,
    output_dir: Path,
    dispatch: pd.DataFrame,
    daily: pd.DataFrame,
    q3_style: bool,
) -> tuple[dict[str, bool], dict[str, float]]:
    template_path = TEMPLATE_DIR / f"{name}.xlsm"
    result_path = output_dir / f"{name}.xlsm"
    template = openpyxl.load_workbook(template_path, read_only=False, data_only=False, keep_vba=True)
    workbook = openpyxl.load_workbook(result_path, read_only=False, data_only=True, keep_vba=True)
    template_sheets = _sheet_snapshot(template)
    result_prefix = _sheet_snapshot(workbook)[:len(template_sheets)]
    checks = {
        "original_sheet_names_order_states": result_prefix == template_sheets,
        "required_added_sheets": workbook.sheetnames[len(template_sheets):] == ["充放电量", "紧急购电量"],
        # 官方文件虽为xlsm，当前包可能没有vbaProject.bin；“保真”应解释为
        # 结果与模板一致：有则哈希相同，无则两边都无，不能把官方无宏误判为失败。
        "vba_preserved": _vba_sha256(template_path) == _vba_sha256(result_path),
    }
    checks["original_headers_unchanged"] = all(
        template[sheet].cell(1, column).value == workbook[sheet].cell(1, column).value
        for sheet in template.sheetnames
        for column in range(1, template[sheet].max_column + 1)
    )
    metrics: dict[str, float] = {}
    tables = [
        sheet for sheet in workbook.worksheets
        if sheet.max_row == 335 and sheet.max_column == 147
    ]
    if q3_style:
        initial = [sheet for sheet in tables if sheet.title == "计划购电量"]
        final = [sheet for sheet in tables if sheet.title == "计划购电量 (3)"]
        checks["table_roles"] = len(initial) == 1 and len(final) == 1
        sheet_specs = [
            (initial[0], "initial_contract", "initial_contract_cost_actual"),
            (final[0], "final_contract", "total_cost_actual"),
        ] if checks["table_roles"] else []
    else:
        visible = [sheet for sheet in tables if sheet.title == "计划购电量 (2)" and sheet.sheet_state == "visible"]
        checks["table_roles"] = len(visible) == 1
        sheet_specs = [(visible[0], "grid_contract", "cost_normal")] if visible else []
        # result4-2中隐藏的“计划购电量”不是题面要求的正式输出，生成器必须保持原值。
        if "计划购电量" in template.sheetnames and "计划购电量" in workbook.sheetnames:
            source = template["计划购电量"]
            target = workbook["计划购电量"]
            unchanged = all(
                source.cell(r, c).value == target.cell(r, c).value
                for r in range(1, source.max_row + 1)
                for c in range(1, source.max_column + 1)
            )
            checks["unused_hidden_plan_unchanged"] = unchanged

    table_error = 0.0
    table_dates_ok = True
    for sheet, quantity_column, cost_column in sheet_specs:
        for row, (day_iso, group) in enumerate(dispatch.groupby("date", sort=True), start=2):
            group = group.sort_values("slot")
            table_dates_ok &= str(sheet.cell(row, 1).value)[:10] == day_iso
            for slot, value in enumerate(group[quantity_column].to_numpy(float)):
                cell = sheet.cell(row, 2 + slot).value
                table_error = max(table_error, abs(float(cell) - float(value)))
            table_error = max(
                table_error,
                abs(float(sheet.cell(row, 146).value) - float(group[quantity_column].sum())),
                abs(float(sheet.cell(row, 147).value) - float(daily.loc[day_iso, cost_column])),
            )
    metrics["workbook_table_max_error"] = table_error
    checks["workbook_table_dates"] = table_dates_ok
    checks["workbook_table_values"] = bool(sheet_specs) and table_error < TOL

    boundary_branch = {"result4-2": "q4_2", "result4-3": "q4_3"}[name]
    boundary_path = output_dir / f"{boundary_branch}_wallclock_boundary.csv"
    if not boundary_path.exists():
        workbook.close()
        template.close()
        raise FileNotFoundError(
            f"缺少{boundary_path.name}，不能独立核验2月1日墙钟边界"
        )
    wallclock_rows = pd.concat(
        [pd.read_csv(boundary_path), dispatch], ignore_index=True, sort=False
    )
    wallclock_rows["interval_start"] = pd.to_datetime(wallclock_rows["interval_start"])
    wallclock_rows["interval_end"] = pd.to_datetime(wallclock_rows["interval_end"])
    cd_error = 0.0
    cd = workbook["充放电量"] if "充放电量" in workbook.sheetnames else None
    if cd is not None:
        for row, day_iso in enumerate(EXPECTED_DAYS, start=2):
            midnight = pd.Timestamp(day_iso)
            next_midnight = midnight + pd.Timedelta(days=1)
            group = wallclock_rows[
                (wallclock_rows.interval_start >= midnight)
                & (wallclock_rows.interval_end <= next_midnight)
            ].sort_values("interval_start")
            expected_starts = list(pd.date_range(midnight, periods=144, freq="10min"))
            if len(group) != 144 or group.interval_start.tolist() != expected_starts:
                raise AssertionError(f"{name} {day_iso} 独立墙钟时间轴不完整")
            for block in range(6):
                left = midnight + pd.Timedelta(hours=4 * block)
                right = left + pd.Timedelta(hours=4)
                part = group[
                    (group.interval_start >= left) & (group.interval_end <= right)
                ]
                cd_error = max(
                    cd_error,
                    abs(float(cd.cell(row, 2 + 2 * block).value) - float(part.charge_actual.sum())),
                    abs(float(cd.cell(row, 3 + 2 * block).value) - float(part.discharge_actual.sum())),
                )
            cd_error = max(
                cd_error,
                abs(float(cd.cell(row, 14).value) - float(group.iloc[0].soc_start)),
                abs(float(cd.cell(row, 15).value) - float(group.iloc[-1].soc_end)),
            )
    metrics["workbook_charge_discharge_max_error"] = cd_error
    checks["workbook_charge_discharge"] = cd is not None and cd.max_row == 335 and cd.max_column == 15 and cd_error < TOL

    expected = _expected_emergency(dispatch)
    actual: list[tuple[str | None, str, float]] = []
    em = workbook["紧急购电量"] if "紧急购电量" in workbook.sheetnames else None
    current_day: str | None = None
    if em is not None:
        for row in range(2, em.max_row + 1):
            raw_day = em.cell(row, 1).value
            if raw_day is not None:
                current_day = str(raw_day)[:10]
            actual.append((current_day, str(em.cell(row, 2).value), float(em.cell(row, 3).value)))
    structure_ok = len(actual) == len(expected) and all(a[:2] == e[:2] for a, e in zip(actual, expected))
    emergency_error = max((abs(a[2] - e[2]) for a, e in zip(actual, expected)), default=0.0) if len(actual) == len(expected) else float("inf")
    metrics["workbook_emergency_max_error"] = emergency_error
    metrics["workbook_emergency_interval_count"] = float(len(expected))
    checks["workbook_emergency"] = structure_ok and emergency_error < TOL
    template.close()
    workbook.close()
    return checks, metrics


def _check_price_vintages(path: Path) -> tuple[dict[str, bool], dict[str, float]]:
    frame = pd.read_csv(
        path,
        parse_dates=["issue_time", "target_time", "source_time", "source_observed_at"],
    )
    required = {
        "issue_time", "target_time", "source_time", "source_observed_at", "price_forecast", "method",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"q4_price_vintages.csv 缺列：{sorted(missing)}")
    checks = {
        "price_vintage_nonempty": not frame.empty,
        "price_vintage_source_causal": bool((frame.source_observed_at <= frame.issue_time).all()),
        "price_vintage_target_future": bool((frame.target_time >= frame.issue_time).all()),
        "price_vintage_source_time_consistent": bool((frame.source_time <= frame.source_observed_at).all()),
        "price_vintage_forecast_finite_nonnegative": bool(
            np.isfinite(frame.price_forecast.to_numpy(float)).all()
            and (frame.price_forecast.to_numpy(float) >= 0).all()
        ),
    }
    lead_days = (frame.target_time - frame.source_time).dt.total_seconds() / 86400.0
    metrics = {
        "price_vintage_rows": float(len(frame)),
        "price_vintage_min_lead_days": float(lead_days.min()),
        "price_vintage_max_lead_days": float(lead_days.max()),
    }
    return checks, metrics


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    prices = _load_actual_prices()
    q42 = _read_dispatch(args.output_dir / "q4_2_dispatch.csv", q3_style=False)
    q43 = _read_dispatch(args.output_dir / "q4_3_dispatch.csv", q3_style=True)
    q42_daily = _read_daily(args.output_dir / "q4_2_daily_summary.csv")
    q43_daily = _read_daily(args.output_dir / "q4_3_daily_summary.csv")

    checks: dict[str, bool] = {}
    metrics: dict[str, float] = {}
    for prefix, frame, q3_style in (("q4_2", q42, False), ("q4_3", q43, True)):
        sub_checks, sub_metrics = _base_dispatch_checks(frame, prices, q3_style)
        checks.update({f"{prefix}_{key}": value for key, value in sub_checks.items()})
        metrics.update({f"{prefix}_{key}": value for key, value in sub_metrics.items()})
    sub_checks, sub_metrics = _check_q4_2_costs(q42, q42_daily)
    checks.update({f"q4_2_{key}": value for key, value in sub_checks.items()})
    metrics.update({f"q4_2_{key}": value for key, value in sub_metrics.items()})
    sub_checks, sub_metrics, _ = _check_q4_3_ledger(
        q43, q43_daily, args.output_dir / "q4_3_versions.csv"
    )
    checks.update({f"q4_3_{key}": value for key, value in sub_checks.items()})
    metrics.update({f"q4_3_{key}": value for key, value in sub_metrics.items()})
    for name, frame, daily, style in (
        ("result4-2", q42, q42_daily, False),
        ("result4-3", q43, q43_daily, True),
    ):
        sub_checks, sub_metrics = _check_workbook(name, args.output_dir, frame, daily, style)
        checks.update({f"{name}_{key}": value for key, value in sub_checks.items()})
        metrics.update({f"{name}_{key}": value for key, value in sub_metrics.items()})
    sub_checks, sub_metrics = _check_price_vintages(args.output_dir / "q4_price_vintages.csv")
    checks.update(sub_checks)
    metrics.update(sub_metrics)

    checks = {key: bool(value) for key, value in checks.items()}
    failed = [key for key, value in checks.items() if not value]
    payload = {
        "summary": {"total": len(checks), "passed": len(checks) - len(failed), "failed": len(failed)},
        "failed_checks": failed,
        "checks": checks,
        "metrics": metrics,
    }
    (args.output_dir / "q4_validation.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
