"""独立读取Q3 CSV与工作簿，复算物理、版本费用和导出一致性。"""

from __future__ import annotations

import json
from pathlib import Path
from datetime import date

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/q3"
TOL = 1e-6


def _time_label(ts: pd.Timestamp, plan_day: str) -> str:
    label = ts.strftime("%H:%M")
    if ts.date() == date.fromisoformat(plan_day) + pd.Timedelta(days=1):
        label += "+1"
    return label


def _expected_emergency_rows(dispatch: pd.DataFrame) -> list[tuple[str, str, float]]:
    """独立按时间连续性合并紧急购电槽，不调用正式导出实现。"""
    expected: list[tuple[str, str, float]] = []
    for day, group in dispatch.groupby("date", sort=True):
        run: list[object] = []
        for row in group.sort_values("slot").itertuples():
            if row.grid_emergency > TOL:
                if run and run[-1].interval_end != row.interval_start:
                    expected.append((day, f"{_time_label(run[0].interval_start, day)}-{_time_label(run[-1].interval_end, day)}", float(sum(x.grid_emergency for x in run))))
                    run = []
                run.append(row)
            elif run:
                expected.append((day, f"{_time_label(run[0].interval_start, day)}-{_time_label(run[-1].interval_end, day)}", float(sum(x.grid_emergency for x in run))))
                run = []
        if run:
            expected.append((day, f"{_time_label(run[0].interval_start, day)}-{_time_label(run[-1].interval_end, day)}", float(sum(x.grid_emergency for x in run))))
    return expected


def main() -> int:
    d = pd.read_csv(OUT / "q3_dispatch.csv", parse_dates=["interval_start", "interval_end"])
    v = pd.read_csv(OUT / "q3_versions.csv", parse_dates=["issue_time"])
    daily = pd.read_csv(OUT / "q3_daily_summary.csv").set_index("date")
    checks = {}
    checks["shape_334x144"] = len(d) == 334*144 and d.groupby("date").size().eq(144).all()
    balance = d.grid_delivered+d.grid_emergency+d.pv_used+d.discharge_actual-d.load_actual-d.charge_actual
    soc_step = d.soc_start+0.9*d.charge_actual-d.discharge_actual/0.9-d.soc_end
    checks["energy_balance"] = float(balance.abs().max()) < 1e-9
    checks["soc_transition"] = float(soc_step.abs().max()) < 1e-9
    checks["soc_bounds"] = d[["soc_start","soc_end"]].min().min() >= 1200-TOL and d[["soc_start","soc_end"]].max().max() <= 10800+TOL
    checks["power_bounds"] = d[["charge_actual","discharge_actual"]].max().max() <= 5000/6+TOL
    checks["no_simultaneous"] = not ((d.charge_actual>TOL)&(d.discharge_actual>TOL)).any()
    checks["contract_identity"] = float((d.final_contract-d.grid_delivered-d.grid_unused).abs().max()) < TOL
    checks["pv_identity"] = float((d.pv_available-d.pv_used-d.pv_curtail).abs().max()) < TOL
    ordered = d.sort_values(["date","slot"])
    checks["slot_order"] = all(g.slot.tolist() == list(range(144)) for _,g in ordered.groupby("date"))
    day_ends = ordered.groupby("date").soc_end.last().iloc[:-1].to_numpy()
    next_starts = ordered.groupby("date").soc_start.first().iloc[1:].to_numpy()
    cross_day_soc_max = float(np.max(np.abs(day_ends-next_starts)))
    checks["cross_day_soc"] = cross_day_soc_max < TOL
    starts = ordered.interval_start.to_numpy()
    ends = ordered.interval_end.to_numpy()
    time_gap_max_seconds = float(np.max(np.abs((starts[1:] - ends[:-1]) / np.timedelta64(1, "s"))))
    checks["interval_continuity"] = time_gap_max_seconds < TOL

    costs = []
    version_ok = True
    final_match = True
    dispatch_by_key = {
        (str(row.date), int(row.slot)): row
        for row in d.itertuples(index=False)
    }
    for (day, slot), group in v.groupby(["date","slot"], sort=False):
        group = group.sort_values("version")
        dispatch_row = dispatch_by_key[(str(day), int(slot))]
        if group.version.tolist() != list(range(len(group))) or (group.issue_time > dispatch_row.interval_start).any():
            version_ok = False
        quantities = group.quantity.to_numpy(float)
        price = float(dispatch_row.price)
        initial = price*quantities[0]
        delta = np.diff(quantities)
        up = 1.5*price*np.maximum(delta,0).sum()
        down = 0.5*price*np.maximum(-delta,0).sum()
        rec = dispatch_row
        emergency = 5*price*float(rec.grid_emergency)
        costs.append((day, initial, up, down, emergency, initial+up-down+emergency))
        final_match &= abs(quantities[0]-rec.initial_contract)<TOL and abs(quantities[-1]-rec.final_contract)<TOL
    checks["version_sequence_and_timing"] = version_ok
    checks["version_dispatch_match"] = bool(final_match)
    c = pd.DataFrame(costs, columns=["date","initial","up","down","emergency","total"]).groupby("date").sum()
    checks["daily_ledger"] = all(float((c[col]-daily[target]).abs().max()) < TOL for col,target in (("initial","initial_contract_cost"),("up","up_cost"),("down","down_credit"),("emergency","emergency_cost"),("total","total_cost")))

    # 普通模式支持 O(1) 随机单元格读取；read_only 下反复 cell() 会退化为重复扫描。
    wb = openpyxl.load_workbook(OUT / "result3.xlsm", read_only=False, data_only=True, keep_vba=True)
    tables = [name for name in wb.sheetnames if wb[name].max_row == 335 and wb[name].max_column == 147]
    checks["workbook_sheets"] = len(tables)==2 and "充放电量" in wb.sheetnames and "紧急购电量" in wb.sheetnames
    excel_max = 0.0; total_max = 0.0; quantity_total_max = 0.0
    cd_max = 0.0; emergency_quantity_max = 0.0
    emergency_structure_ok = False
    if len(tables)==2:
        pws,aws=wb[tables[0]],wb[tables[1]]
        for i,(day,g) in enumerate(ordered.groupby("date"),2):
            g=g.sort_values("slot")
            for j,row in enumerate(g.itertuples(),2):
                excel_max=max(excel_max,abs(float(pws.cell(i,j).value)-row.initial_contract),abs(float(aws.cell(i,j).value)-row.final_contract))
            quantity_total_max=max(quantity_total_max,
                                   abs(float(pws.cell(i,146).value)-float(g.initial_contract.sum())),
                                   abs(float(aws.cell(i,146).value)-float(g.final_contract.sum())))
            total_max=max(total_max,abs(float(pws.cell(i,147).value)-daily.loc[day,"initial_contract_cost"]),abs(float(aws.cell(i,147).value)-daily.loc[day,"total_cost"]))
        cdws = wb["充放电量"]
        for i, (day, g) in enumerate(ordered.groupby("date"), 2):
            g = g.sort_values("slot")
            for block in range(6):
                block_rows = g[(g.slot >= 24*block) & (g.slot < 24*(block+1))]
                cd_max = max(cd_max,
                             abs(float(cdws.cell(i, 2+2*block).value)-float(block_rows.charge_actual.sum())),
                             abs(float(cdws.cell(i, 3+2*block).value)-float(block_rows.discharge_actual.sum())))
            cd_max = max(cd_max,
                         abs(float(cdws.cell(i,14).value)-float(g.iloc[0].soc_start)),
                         abs(float(cdws.cell(i,15).value)-float(g.iloc[-1].soc_end)))

        expected_emergency = _expected_emergency_rows(ordered)
        emws = wb["紧急购电量"]
        actual_emergency = []
        inherited_day = None
        for i in range(2, emws.max_row+1):
            raw_day = emws.cell(i,1).value
            if raw_day is not None:
                inherited_day = str(raw_day)[:10]
            actual_emergency.append((inherited_day, str(emws.cell(i,2).value), float(emws.cell(i,3).value)))
        emergency_structure_ok = len(actual_emergency) == len(expected_emergency) and all(
            a[0] == e[0] and a[1] == e[1] for a,e in zip(actual_emergency, expected_emergency)
        )
        if len(actual_emergency) == len(expected_emergency):
            emergency_quantity_max = max((abs(a[2]-e[2]) for a,e in zip(actual_emergency, expected_emergency)), default=0.0)
    checks["workbook_values"] = excel_max<TOL and total_max<TOL and quantity_total_max<TOL
    checks["workbook_charge_discharge"] = cd_max<TOL
    checks["workbook_emergency_intervals"] = emergency_structure_ok and emergency_quantity_max<TOL
    wb.close()

    summary_file=json.loads((OUT/"q3_replay_summary.json").read_text(encoding="utf-8"))
    checks["summary_total"] = abs(summary_file["total_cost"]-daily.total_cost.sum())<TOL
    checks = {name: bool(value) for name, value in checks.items()}
    failed=[k for k,vv in checks.items() if not vv]
    payload={
        "summary":{"total":len(checks),"passed":len(checks)-len(failed),"failed":len(failed)},
        "checks":checks,
        "metrics":{"max_energy_residual":float(balance.abs().max()),"max_soc_residual":float(soc_step.abs().max()),
                   "max_cross_day_soc_error":cross_day_soc_max,"max_interval_gap_seconds":time_gap_max_seconds,
                   "workbook_cell_max_error":excel_max,"workbook_quantity_total_max_error":quantity_total_max,
                   "workbook_cost_total_max_error":total_max,"workbook_charge_discharge_max_error":cd_max,
                   "workbook_emergency_quantity_max_error":emergency_quantity_max,
                   "workbook_emergency_interval_count":len(_expected_emergency_rows(ordered)),
                   "total_cost":float(daily.total_cost.sum())},
    }
    (OUT/"q3_validation.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(payload,ensure_ascii=False,indent=2))
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
