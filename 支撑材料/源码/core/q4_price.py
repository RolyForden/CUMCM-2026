"""第四问的实时电价读取与严格因果预测。

优化器只能调用 :func:`forecast_prices_as_of`；实际价格仅供执行器和
独立核算器使用。多日视野中的全部目标槽都统一受传入的 decision_time
约束，不能把未来目标日伪装成新的决策时刻。
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from numbers import Real
from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl
import pandas as pd

from core.slot_adapter import N_SLOTS, build_day_slots


ROOT = Path(__file__).resolve().parents[2]
ATTACHMENT4 = ROOT / "data/raw/official/附件4.xlsm"
ATTACHMENT1 = ROOT / "data/raw/official/附件1.xlsm"

OUTPUT_COLUMNS = [
    "issue_time",
    "target_time",
    "source_time",
    "source_observed_at",
    "price_forecast",
    "method",
]


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return pd.to_datetime(str(value), errors="raise").date()


def load_attachment4_prices(path: str | Path = ATTACHMENT4) -> pd.DataFrame:
    """一次读取附件4，返回2025年365×144的长表。"""
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1 (3)"]
    rows: list[dict] = []
    seen_dates: set[date] = set()
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if excel_row[0] in (None, ""):
            continue
        day = _as_date(excel_row[0])
        if day in seen_dates:
            wb.close()
            raise ValueError(f"附件4日期重复: {day}")
        seen_dates.add(day)
        values = excel_row[1:145]
        if len(values) != N_SLOTS or any(
            not isinstance(value, Real) or isinstance(value, bool) for value in values
        ):
            wb.close()
            raise ValueError(f"附件4日期 {day} 必须恰含144个数值")
        for slot, value in zip(build_day_slots(day), values):
            rows.append(
                {
                    "date": day,
                    "slot": slot.slot_id,
                    "target_time": slot.interval_start,
                    "observed_at": slot.interval_end,
                    "price_actual": float(value),
                }
            )
    wb.close()
    out = pd.DataFrame(rows)
    expected_dates = {date(2025, 1, 1) + timedelta(days=k) for k in range(365)}
    if seen_dates != expected_dates or len(out) != 365 * N_SLOTS:
        raise ValueError(
            f"附件4必须覆盖2025年365天且每天144槽，当前日期数={len(seen_dates)}，记录数={len(out)}"
        )
    if out.duplicated("target_time").any() or out["price_actual"].isna().any():
        raise ValueError("附件4长表存在重复目标时刻或空价格")
    return out


def load_attachment1_price_prior(path: str | Path = ATTACHMENT1) -> pd.DataFrame:
    """读取附件1典型日价格，作为一月冷启动和历史缺失回退。"""
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1 (2)"]
    rows = list(ws.iter_rows(min_row=2, max_row=145, values_only=True))
    wb.close()
    values = [row[1] for row in rows]
    if len(values) != N_SLOTS or any(
        not isinstance(value, Real) or isinstance(value, bool) for value in values
    ):
        raise ValueError("附件1典型日价格必须恰含144个数值")
    return pd.DataFrame({"slot": range(N_SLOTS), "price_prior": map(float, values)})


def _slot_id(target_time: datetime) -> int:
    if target_time.hour == 0 and target_time.minute == 0:
        return 143
    minutes = target_time.hour * 60 + target_time.minute
    slot = minutes // 10 - 1
    if target_time.second or target_time.microsecond or minutes % 10 or not 0 <= slot < 143:
        raise ValueError(f"目标时刻不是方案A的槽位起点: {target_time!r}")
    return slot


def _normalise_method(method: str) -> str:
    key = method.strip().lower().replace("_", "-")
    aliases = {
        "d-1": "D-1",
        "d1": "D-1",
        "d-7": "D-7",
        "d7": "D-7",
        "same-weekday": "same-weekday-weighted",
        "same-weekday-weighted": "same-weekday-weighted",
        "weekday-weighted": "same-weekday-weighted",
    }
    if key not in aliases:
        raise ValueError(f"未知电价预测方法: {method}")
    return aliases[key]


def _validate_actual_prices(actual_prices: pd.DataFrame) -> pd.DataFrame:
    required = {"target_time", "observed_at", "price_actual"}
    missing = required - set(actual_prices.columns)
    if missing:
        raise ValueError(f"实际电价缺少字段: {sorted(missing)}")
    actual = actual_prices.copy()
    actual["target_time"] = pd.to_datetime(actual["target_time"])
    actual["observed_at"] = pd.to_datetime(actual["observed_at"])
    if actual.empty or actual["target_time"].duplicated().any():
        raise ValueError("实际电价不能为空，且target_time必须唯一")
    if actual[["target_time", "observed_at", "price_actual"]].isna().any().any():
        raise ValueError("实际电价含空值")
    if (actual["observed_at"] <= actual["target_time"]).any():
        raise ValueError("实际电价的observed_at必须晚于目标槽起点")
    return actual.sort_values("target_time").reset_index(drop=True)


def _prior_by_slot() -> dict[int, float]:
    prior = load_attachment1_price_prior()
    return dict(zip(prior["slot"].astype(int), prior["price_prior"].astype(float)))


def _prior_row(target: datetime, decision_time: datetime, label: str, prior: dict[int, float]) -> dict:
    return {
        "issue_time": decision_time,
        "target_time": target,
        "source_time": pd.NaT,
        "source_observed_at": decision_time,
        "price_forecast": prior[_slot_id(target)],
        "method": label,
    }


def forecast_prices_as_of(
    actual_prices: pd.DataFrame,
    decision_time: datetime,
    target_times: Iterable[datetime],
    method: str,
) -> pd.DataFrame:
    """在同一真实决策时刻，为任意多日视野生成因果电价预测。

    ``D-1`` 严格要求一天前同刻已经观测完成，不做隐式回退；因此午夜
    决策的方案A末槽会被拒绝。``D-7`` 在七天前来源尚不可用或缺失时，
    依次尝试D-14、D-21等同星期来源，仍无来源才回退附件1典型日先验。
    一月1日至7日统一执行冻结的冷启动规则，不因未来目标日期而改变。
    """
    decision_time = pd.Timestamp(decision_time).to_pydatetime()
    canonical = _normalise_method(method)
    actual = _validate_actual_prices(actual_prices)
    lookup = actual.set_index("target_time")
    targets = [pd.Timestamp(value).to_pydatetime() for value in target_times]
    if not targets:
        raise ValueError("target_times不能为空")
    if len(set(targets)) != len(targets):
        raise ValueError("target_times不得重复")
    prior = _prior_by_slot()
    available = actual[actual["observed_at"] <= decision_time]

    rows: list[dict] = []
    cold_start = decision_time.date() <= date(2025, 1, 7)
    for target in targets:
        slot = _slot_id(target)
        if cold_start:
            slot_history = available[
                available["target_time"].map(lambda value: _slot_id(value.to_pydatetime())) == slot
            ]
            if slot_history.empty:
                rows.append(_prior_row(target, decision_time, "cold-start:attachment1-prior", prior))
            else:
                latest = slot_history.iloc[-1]
                rows.append(
                    {
                        "issue_time": decision_time,
                        "target_time": target,
                        "source_time": latest.target_time,
                        "source_observed_at": slot_history["observed_at"].max(),
                        "price_forecast": float(slot_history["price_actual"].mean()),
                        "method": "cold-start:observed-slot-mean",
                    }
                )
            continue

        if canonical == "D-1":
            source_time = pd.Timestamp(target - timedelta(days=1))
            if source_time not in lookup.index:
                raise ValueError(f"{target.isoformat()} 缺少D-1实际电价")
            source = lookup.loc[source_time]
            if source.observed_at > decision_time:
                raise ValueError(f"{target.isoformat()} 的D-1来源在决策时刻尚未观测完成")
            rows.append(
                {
                    "issue_time": decision_time,
                    "target_time": target,
                    "source_time": source_time,
                    "source_observed_at": source.observed_at,
                    "price_forecast": float(source.price_actual),
                    "method": "D-1",
                }
            )
            continue

        if canonical == "D-7":
            chosen = None
            chosen_source_time = None
            chosen_lag = None
            for lag in range(7, 7 * 53, 7):
                source_time = pd.Timestamp(target - timedelta(days=lag))
                if source_time in lookup.index:
                    source = lookup.loc[source_time]
                    if source.observed_at <= decision_time:
                        chosen = source
                        chosen_source_time = source_time
                        chosen_lag = lag
                        break
            if chosen is None:
                rows.append(_prior_row(target, decision_time, "D-7:fallback-prior", prior))
            else:
                rows.append(
                    {
                        "issue_time": decision_time,
                        "target_time": target,
                        "source_time": chosen_source_time,
                        "source_observed_at": chosen.observed_at,
                        "price_forecast": float(chosen.price_actual),
                        "method": "D-7" if chosen_lag == 7 else f"D-7:fallback-D-{chosen_lag}",
                    }
                )
            continue

        sources = []
        for lag in (7, 14, 21, 28):
            source_time = pd.Timestamp(target - timedelta(days=lag))
            if source_time in lookup.index:
                source = lookup.loc[source_time]
                if source.observed_at <= decision_time:
                    sources.append((lag, source_time, source))
        if not sources:
            rows.append(_prior_row(target, decision_time, "same-weekday-weighted:fallback-prior", prior))
        else:
            weights = np.array([0.8 ** (lag // 7 - 1) for lag, _, _ in sources], dtype=float)
            weights /= weights.sum()
            latest_time, latest = max(
                ((source_time, source) for _, source_time, source in sources),
                key=lambda item: item[0],
            )
            rows.append(
                {
                    "issue_time": decision_time,
                    "target_time": target,
                    "source_time": latest_time,
                    "source_observed_at": max(source.observed_at for _, _, source in sources),
                    "price_forecast": float(
                        sum(
                            weight * source.price_actual
                            for weight, (_, _, source) in zip(weights, sources)
                        )
                    ),
                    "method": f"same-weekday-weighted:{len(sources)}",
                }
            )

    out = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    if (out["issue_time"] != decision_time).any():
        raise AssertionError("多日视野未统一使用传入的真实决策时刻")
    if (out["source_observed_at"] > out["issue_time"]).any():
        raise AssertionError("电价预测使用了决策时刻之后才观测到的来源")
    if out["price_forecast"].isna().any() or (out["price_forecast"] < 0).any():
        raise ValueError("电价预测含空值或负值")
    return out
