"""从已完成的Q3 CSV生成官方 result3.xlsm，不重新运行全年模型。"""

from __future__ import annotations

from datetime import date
import hashlib
import json
from pathlib import Path
import shutil
import sys

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from generate_q2_result import merge_emergency_intervals
from core.wallclock_output import aggregate_wallclock_days

OUT = ROOT / "outputs/q3"
TEMPLATE = ROOT / "data/raw/official/附件5/result3.xlsm"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024*1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> int:
    dispatch = pd.read_csv(OUT / "q3_dispatch.csv")
    daily = pd.read_csv(OUT / "q3_daily_summary.csv").set_index("date")
    boundary_path = OUT / "q3_wallclock_boundary.csv"
    if not boundary_path.exists():
        raise FileNotFoundError(
            "缺少q3_wallclock_boundary.csv：现有正式CSV无法可靠恢复2月1日00:00边界，"
            "拒绝猜测或沿用错位分块"
        )
    boundary = pd.read_csv(boundary_path)
    days = sorted(dispatch.date.unique())
    wallclock = {
        summary.day.isoformat(): summary
        for summary in aggregate_wallclock_days(
            dispatch.to_dict("records"),
            [date.fromisoformat(str(day)) for day in days],
            boundary_records=boundary.to_dict("records"),
        )
    }
    target = OUT / "result3.xlsm"
    shutil.copy2(TEMPLATE, target)
    wb = openpyxl.load_workbook(target, keep_vba=True)
    tables = [name for name in wb.sheetnames if wb[name].max_row == 335 and wb[name].max_column == 147]
    if len(tables) != 2:
        wb.close(); raise ValueError(f"result3模板的计划/调整工作表数量={len(tables)}")
    plan_ws, adjusted_ws = (wb[tables[0]], wb[tables[1]])
    if len(days) != 334:
        wb.close(); raise ValueError("Q3正式记录不是334天")
    by_day: dict[str, list[dict]] = {}
    for row in dispatch.to_dict("records"):
        by_day.setdefault(row["date"], []).append(row)
    interval_rows = []
    for i, day_iso in enumerate(days, start=2):
        rows = sorted(by_day[day_iso], key=lambda r: r["slot"])
        if len(rows) != 144 or str(plan_ws.cell(i, 1).value)[:10] != day_iso:
            wb.close(); raise ValueError(f"{day_iso}模板日期或槽数错误")
        for row in rows:
            col = 2 + int(row["slot"])
            plan_ws.cell(i, col, float(row["initial_contract"]))
            adjusted_ws.cell(i, col, float(row["final_contract"]))
        plan_ws.cell(i, 146, float(sum(r["initial_contract"] for r in rows)))
        plan_ws.cell(i, 147, float(daily.loc[day_iso, "initial_contract_cost"]))
        adjusted_ws.cell(i, 146, float(sum(r["final_contract"] for r in rows)))
        adjusted_ws.cell(i, 147, float(daily.loc[day_iso, "total_cost"]))
        compact = [
            {**r, "charge": r["charge_actual"], "discharge": r["discharge_actual"],
             "grid_emergency": r["grid_emergency"]}
            for r in rows
        ]
        for item in merge_emergency_intervals(date.fromisoformat(day_iso), compact):
            interval_rows.append({"date": day_iso, **item})

    cd = wb.create_sheet("充放电量")
    block_labels = ("0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00")
    header = ["日期"] + [x for label in block_labels for x in (f"{label}充电量", f"{label}放电量")] + ["0:00储电量", "24:00储电量"]
    for c, value in enumerate(header, 1): cd.cell(1, c, value)
    for i, day_iso in enumerate(days, 2):
        summary = wallclock[str(day_iso)]
        cd.cell(i, 1, day_iso)
        for b, block in enumerate(summary.blocks):
            cd.cell(i, 2+2*b, block.charge); cd.cell(i, 3+2*b, block.discharge)
        cd.cell(i, 14, summary.soc_0000); cd.cell(i, 15, summary.soc_2400)

    em = wb.create_sheet("紧急购电量")
    for c, value in enumerate(("日期", "紧急购电时间段", "紧急购电量"), 1): em.cell(1, c, value)
    row_no = 2
    current_day = None
    for item in interval_rows:
        em.cell(row_no, 1, item["date"] if item["date"] != current_day else None)
        em.cell(row_no, 2, f"{item['start']}-{item['end_label']}")
        em.cell(row_no, 3, float(item["quantity"]))
        current_day = item["date"]; row_no += 1
    wb.save(target); wb.close()

    audit = {
        "template_sha256": sha256(TEMPLATE), "result3_sha256": sha256(target),
        "days": len(days), "plan_sheet": tables[0], "adjusted_sheet": tables[1],
        "emergency_intervals": len(interval_rows), "sheets_added": ["充放电量", "紧急购电量"],
    }
    (OUT / "q3_workbook_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
