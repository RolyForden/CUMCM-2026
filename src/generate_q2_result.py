"""按 D011（rolling-7d）生成并审计第二问正式结果。

产物（outputs/q2/）：
- result2.xlsm：官方模板复制后原位填写“计划购电量”，按题面补齐“充放电量”
  与“紧急购电量”工作表；官方原有工作表、顺序、隐藏状态与 VBA 容器保持。
- q2_dispatch.csv：逐槽执行记录（评价窗口内）。
- q2_daily_summary.csv：逐日费用、紧急购电、库存与削减汇总。
- q2_emergency_intervals.csv：合并后的紧急购电时间段。
- q2_audit.json：生成期审计与重开回读核验。

费用一律由独立核算器复算；本脚本不修改 data/raw。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import date, datetime, timedelta
from pathlib import Path
import sys
from typing import Any

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from core.q2_replay import (
    Strategy,
    Terminal,
    daily_price,
    load_window_actuals,
    prior_from_attachment1,
    replay,
)
from core.slot_adapter import build_day_slots

DEFAULT_TEMPLATE = ROOT / "data/raw/official/附件5/result2.xlsm"
DEFAULT_OUTPUT = ROOT / "outputs/q2"
START = date(2025, 1, 1)
EVAL_START = date(2025, 2, 1)
EVAL_END = date(2025, 12, 31)
CACHE = ROOT / "experiments" / "q2_actuals_cache.csv"
EMERGENCY_EPS = 1e-6

# 表 2 / result2 充放电量：6 个 4 小时块，每块 24 槽（沿用 template_io 口径）
CD_BLOCK_LABELS = (
    "0:00-4:00", "4:00-8:00", "8:00-12:00",
    "12:00-16:00", "16:00-20:00", "20:00-24:00",
)
CD_BLOCK_SLOTS = 24


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_actuals() -> "pd.DataFrame":
    import pandas as pd
    if CACHE.exists():
        cached = pd.read_csv(CACHE, parse_dates=["valid_time", "observed_at", "date"])
        cached["date"] = cached["date"].dt.date
        if cached["date"].min() == START and cached["date"].max() == EVAL_END:
            return cached
    actuals = load_window_actuals(START, EVAL_END)
    actuals.to_csv(CACHE, index=False)
    return actuals


def _fmt_hhmm(dt: datetime) -> str:
    return f"{dt.hour:02d}:{dt.minute:02d}"


def merge_emergency_intervals(
    day: date, records: list[dict]
) -> list[dict[str, Any]]:
    """把某自然日内相邻的非零紧急购电槽合并为连续时间段。

    records 必须按槽位升序、连续；相邻 = 前槽 interval_end == 后槽 interval_start。
    跨午夜：槽 142 终点为 24:00，槽 143（次日00:00-00:10）终点写作 '00:10+1'。
    """
    slots = build_day_slots(day)
    ordered = sorted(records, key=lambda r: r["slot"])
    merged: list[dict[str, Any]] = []
    run: list[dict] = []
    for rec in ordered:
        if rec["grid_emergency"] > EMERGENCY_EPS:
            if run and run[-1]["interval_end"] != rec["interval_start"]:
                merged.append(_close_run(run))
                run = []
            run.append(rec)
        else:
            if run:
                merged.append(_close_run(run))
                run = []
    if run:
        merged.append(_close_run(run))
    return merged


def _close_run(run: list[dict]) -> dict[str, Any]:
    first_start = datetime.fromisoformat(run[0]["interval_start"])
    last_end = datetime.fromisoformat(run[-1]["interval_end"])
    start_day = first_start.date()
    end_label = _fmt_hhmm(last_end)
    if last_end.date() != start_day:
        end_label += "+1"
    return {
        "start": _fmt_hhmm(first_start),
        "end_label": end_label,
        "quantity": float(sum(r["grid_emergency"] for r in run)),
        "slot_first": run[0]["slot"],
        "slot_last": run[-1]["slot"],
    }


def build_cd_blocks(records: list[dict]) -> tuple[list[float], list[float]]:
    """按 6 个 4 小时块汇总某日的充电量与放电量（每块 24 槽）。"""
    by_slot = {r["slot"]: r for r in records}
    charges, discharges = [], []
    for b in range(6):
        lo = b * CD_BLOCK_SLOTS
        hi = lo + CD_BLOCK_SLOTS
        charges.append(float(sum(by_slot[s]["charge"] for s in range(lo, hi))))
        discharges.append(float(sum(by_slot[s]["discharge"] for s in range(lo, hi))))
    return charges, discharges


def write_result2(template: Path, output: Path, records: list[dict], daily_normal: dict[str, float]) -> None:
    keep_vba = str(template).lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(template, keep_vba=keep_vba)

    plan_sheet = None
    for name in wb.sheetnames:
        if name.startswith("计划购电量"):
            plan_sheet = name
            break
    if plan_sheet is None:
        wb.close()
        raise ValueError("官方 result2 模板缺少计划购电量工作表")
    ws = wb[plan_sheet]

    # 每天 144 槽 + 全天购电量 + 全天购电费
    by_day: dict[str, list[dict]] = {}
    for rec in records:
        by_day.setdefault(rec["date"], []).append(rec)
    days = sorted(by_day)
    expected_days = [
        (EVAL_START + timedelta(days=i)).isoformat()
        for i in range((EVAL_END - EVAL_START).days + 1)
    ]
    if days != expected_days:
        wb.close()
        raise ValueError("执行记录未覆盖 2025-02-01 至 12-31 全部 334 天")

    for i, day_iso in enumerate(days):
        row = i + 2
        day_records = sorted(by_day[day_iso], key=lambda r: r["slot"])
        if len(day_records) != 144:
            wb.close()
            raise ValueError(f"{day_iso} 不是 144 槽")
        label = ws.cell(row=row, column=1).value
        if label is None or str(label)[:10] != day_iso:
            wb.close()
            raise ValueError(f"{day_iso} 模板日期行错位: {label!r}")
        for rec in day_records:
            ws.cell(row=row, column=2 + rec["slot"], value=float(rec["grid_contract"]))
        # 全天购电量 = 当日合同购电之和；全天购电费 = 计划购电量费用（正常电费，
        # 按完成教学口径不含紧急购电费；紧急购电在单独工作表列示）。
        ws.cell(row=row, column=146, value=float(sum(r["grid_contract"] for r in day_records)))
        ws.cell(row=row, column=147, value=float(daily_normal[day_iso]))

    # 充放电量：每天 6 块 + 0:00/24:00 储电量
    cd = wb.create_sheet("充放电量")
    header = ["日期"]
    for label in CD_BLOCK_LABELS:
        header += [f"{label}充电量", f"{label}放电量"]
    header += ["0:00储电量", "24:00储电量"]
    for c, text in enumerate(header, start=1):
        cd.cell(row=1, column=c, value=text)
    for i, day_iso in enumerate(days):
        day_records = sorted(by_day[day_iso], key=lambda r: r["slot"])
        charges, discharges = build_cd_blocks(day_records)
        cd.cell(row=i + 2, column=1, value=day_iso)
        for b in range(6):
            cd.cell(row=i + 2, column=2 + 2 * b, value=charges[b])
            cd.cell(row=i + 2, column=3 + 2 * b, value=discharges[b])
        cd.cell(row=i + 2, column=14, value=float(day_records[0]["soc_start"]))
        cd.cell(row=i + 2, column=15, value=float(day_records[-1]["soc_end"]))

    # 紧急购电量：日期 | 时间段 | 购电量；同一天多段时日期只写首行
    em = wb.create_sheet("紧急购电量")
    for c, text in enumerate(["日期", "紧急购电时间段", "紧急购电量"], start=1):
        em.cell(row=1, column=c, value=text)
    out_row = 2
    for day_iso in days:
        day_records = sorted(by_day[day_iso], key=lambda r: r["slot"])
        intervals = merge_emergency_intervals(date.fromisoformat(day_iso), day_records)
        for j, item in enumerate(intervals):
            if j == 0:
                em.cell(row=out_row, column=1, value=day_iso)
            em.cell(row=out_row, column=2, value=f"{item['start']}-{item['end_label']}")
            em.cell(row=out_row, column=3, value=item["quantity"])
            out_row += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()


def read_back(output: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(output, read_only=True, data_only=True, keep_vba=True)
    ws = None
    for name in wb.sheetnames:
        if name.startswith("计划购电量"):
            ws = wb[name]
            break
    plan_rows = 0
    plan_min = float("inf")
    plan_max = float("-inf")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        plan_rows += 1
        for v in row[1:145]:
            if v is not None:
                plan_min = min(plan_min, float(v))
                plan_max = max(plan_max, float(v))
    cd = wb["充放电量"]
    em = wb["紧急购电量"]
    sheets = {name: {"state": wb[name].sheet_state} for name in wb.sheetnames}
    vba_names = sorted(
        name for name in (wb.vba_archive.namelist() if wb.vba_archive else [])
        if "vba" in name.lower() or "macro" in name.lower()
    )
    wb.close()
    return {
        "plan_rows": plan_rows,
        "plan_min": plan_min,
        "plan_max": plan_max,
        "cd_rows": cd.max_row,
        "cd_cols": cd.max_column,
        "em_rows": em.max_row,
        "sheets": sheets,
        "vba_names": vba_names,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--method", default="same-weekday-4-decay-0.8")
    parser.add_argument("--level", default="point")
    parser.add_argument("--terminal", default="rolling")
    parser.add_argument("--terminal-param", type=float, default=7.0)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    actuals = load_actuals()
    prior = prior_from_attachment1(START)
    price = daily_price()
    terminal = Terminal(args.terminal, args.terminal_param)
    strategy = Strategy(args.method, args.level)

    row = replay(
        strategy, terminal, actuals, prior, price, START, EVAL_START, EVAL_END,
        collect_records=True,
    )
    if not row["all_audits_ok"]:
        raise AssertionError("全年回放存在核算失败日")

    records = row["records"]
    if len(records) != 334 * 144:
        raise AssertionError(f"执行记录槽数 {len(records)} != {334*144}")

    # 逐日费用并入记录，供 summary/费用列使用
    daily_normal = {d["date"]: d["cost_normal"] for d in row["daily"]}

    output = args.output
    result2 = output / "result2.xlsm"
    write_result2(args.template, result2, records, daily_normal)
    audit = read_back(result2)

    # 逐槽 CSV
    with (output / "q2_dispatch.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=["date", "slot", "grid_contract", "grid_emergency",
                        "charge", "discharge", "soc_start", "soc_end",
                        "interval_start", "interval_end"],
        )
        writer.writeheader()
        for rec in records:
            writer.writerow({k: rec[k] for k in writer.fieldnames})

    # 逐日汇总
    with (output / "q2_daily_summary.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(row["daily"][0].keys()))
        writer.writeheader()
        for day_row in row["daily"]:
            if day_row["date"] >= EVAL_START.isoformat():
                writer.writerow(day_row)

    # 紧急购电区间
    interval_rows = []
    for day_row in row["daily"]:
        if day_row["date"] < EVAL_START.isoformat():
            continue
        day_d = date.fromisoformat(day_row["date"])
        day_records = [r for r in records if r["date"] == day_row["date"]]
        for item in merge_emergency_intervals(day_d, day_records):
            interval_rows.append(
                {
                    "date": day_row["date"],
                    "interval": f"{item['start']}-{item['end_label']}",
                    "quantity_kwh": item["quantity"],
                }
            )
    with (output / "q2_emergency_intervals.csv").open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=["date", "interval", "quantity_kwh"])
        writer.writeheader()
        writer.writerows(interval_rows)

    summary = {
        "method": args.method,
        "level": args.level,
        "terminal": terminal.label(),
        "template_sha256": sha256(args.template),
        "result2_sha256": sha256(result2),
        "evaluation_days": len(row["daily"]) - 31,
        "cost_normal": row["cost_normal"],
        "cost_emergency": row["cost_emergency"],
        "cost_total": row["cost_total"],
        "grid_emergency_kwh": row["grid_emergency_kwh"],
        "grid_unused_kwh": row["grid_unused_kwh"],
        "infeasible_days": row["infeasible_days"],
        "all_audits_ok": row["all_audits_ok"],
        "emergency_intervals": len(interval_rows),
        "readback": audit,
    }
    (output / "q2_audit.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
