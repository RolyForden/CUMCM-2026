"""整改后独立交叉验证。

这些反例不调用 ``probe_timing.main``，关键预期值均在本文件硬编码。
默认只打印；传入 ``--output`` 可保存逐项 JSON 证据。
"""

from __future__ import annotations

import argparse
import copy
import json
import shutil
import tempfile
from dataclasses import asdict, replace
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from scipy.optimize import Bounds, LinearConstraint, milp

from core import accountant, data_io, executor, lp_kernel, template_io
from core.params import BatteryParams
from core.slot_adapter import build_day_slots


RESULTS: list[dict[str, object]] = []


def check(name: str, condition: bool, expected: object, actual: object) -> None:
    RESULTS.append(
        {"name": name, "passed": bool(condition), "expected": str(expected), "actual": str(actual)}
    )


def expect_value_error(name: str, fn) -> None:
    try:
        fn()
    except ValueError as exc:
        check(name, True, "ValueError", type(exc).__name__)
    else:
        check(name, False, "ValueError", "no exception")


def make_records() -> list[executor.ExecRecord]:
    slots = build_day_slots(date(2025, 1, 1))[:3]
    return executor.execute_q1(
        grid_contract=np.array([10.0, 10.0, 10.0]),
        price=np.ones(3), load=np.array([10.0, 10.0, 10.0]),
        pv_available=np.zeros(3), pv_used=np.zeros(3), pv_curtail=np.zeros(3),
        charge=np.zeros(3), discharge=np.zeros(3), soc0=5000.0,
        starts=[s.interval_start for s in slots], ends=[s.interval_end for s in slots],
        plan_issue_time=datetime(2025, 1, 1),
    )


def test_efficiency_and_lexicographic() -> None:
    params = BatteryParams()
    slots = build_day_slots(date(2025, 1, 1))[:2]
    recs = executor.execute_q1(
        grid_contract=np.array([100.0, 0.0]), price=np.ones(2),
        load=np.array([0.0, 81.0]), pv_available=np.zeros(2),
        pv_used=np.zeros(2), pv_curtail=np.zeros(2),
        charge=np.array([100.0, 0.0]), discharge=np.array([0.0, 81.0]),
        soc0=5000.0, starts=[s.interval_start for s in slots],
        ends=[s.interval_end for s in slots], plan_issue_time=datetime(2025, 1, 1),
        params=params,
    )
    check("效率-执行器充电100", abs(recs[0].soc_end - 5090.0) < 1e-9, 5090, recs[0].soc_end)
    check("效率-执行器再放电81", abs(recs[1].soc_end - 5000.0) < 1e-9, 5000, recs[1].soc_end)
    aud = accountant.audit(recs, params=params)
    check("效率-独立核算器", aud.ok, True, aud.violations)
    lp_eff = lp_kernel.solve_lp(
        lp_kernel.LpInputs(
            np.array([0.0, 1.0]), np.array([0.0, 81.0]), np.zeros(2),
            5000.0, 5000.0, params=params,
        )
    )
    lp_eff_actual = (lp_eff.charge[0], lp_eff.soc[1], lp_eff.discharge[1], lp_eff.soc[2])
    check(
        "效率-LP状态100到90到81",
        np.allclose(lp_eff_actual, (100.0, 5090.0, 81.0, 5000.0), atol=1e-6),
        (100.0, 5090.0, 81.0, 5000.0), lp_eff_actual,
    )

    zero = lp_kernel.solve_lp(lp_kernel.LpInputs(np.zeros(1), np.zeros(1), np.zeros(1), 5000, 5000))
    check("字典序-零场景费用", zero.cost == 0.0, 0.0, zero.cost)
    check("字典序-零场景最小吞吐", float(zero.charge[0] + zero.discharge[0]) < 1e-10, 0.0, zero.charge[0] + zero.discharge[0])
    one = lp_kernel.solve_lp(lp_kernel.LpInputs(np.array([0.5]), np.array([100.0]), np.zeros(1), 5000, 5000))
    check("字典序-单槽费用不加1e-6", abs(one.cost - 50.0) < 1e-8, 50.0, one.cost)
    check("字典序-第二阶段费用上界", one.cost <= one.primary_optimum + lp_kernel.LEX_TOL + 1e-10,
          f"<= {one.primary_optimum + lp_kernel.LEX_TOL}", one.cost)


def test_dates_and_causality() -> None:
    d1, d2 = date(2025, 1, 1), date(2025, 1, 2)
    jan1 = data_io.attachment2_load(d1)
    jan2 = data_io.attachment2_load(d2)
    check("日期-1月1日与1月2日不同", not np.allclose(jan1.value, jan2.value), "different", "different" if not np.allclose(jan1.value, jan2.value) else "same")

    wb = openpyxl.load_workbook("data/raw/official/附件2.xlsm", read_only=True, data_only=True)
    ws = wb["小区负载"]
    raw = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
    wb.close()
    indices = (0, 71, 143)
    actual = [float(jan2.iloc[k].value) for k in indices]
    expected = [float(raw[k + 1]) for k in indices]
    check("日期-1月2日人工三槽抽查", np.allclose(actual, expected), expected, actual)
    expect_value_error("日期-不存在日期拒绝", lambda: data_io.attachment2_load(date(2024, 12, 31)))

    with tempfile.TemporaryDirectory(prefix="cumcm_date_") as tmp:
        path = Path(tmp) / "duplicate.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws.append(["date"] + list(range(144)))
        row = [datetime(2025, 1, 1)] + [float(k) for k in range(144)]
        ws.append(row); ws.append(row)
        wb.save(path); wb.close()
        expect_value_error("日期-重复日期拒绝", lambda: data_io._read_wide(path, "S", d1))
        extra_path = Path(tmp) / "extra.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S"
        ws.append(["date"] + list(range(145)))
        ws.append([datetime(2025, 1, 1)] + [float(k) for k in range(145)])
        wb.save(extra_path); wb.close()
        expect_value_error("日期-多于144点拒绝", lambda: data_io._read_wide(extra_path, "S", d1))

    decision = datetime(2025, 2, 1, 0, 0)
    hist = data_io.history_actuals_as_of(decision)
    check("防泄漏-一月历史进入", not hist.empty and hist.date.min() == d1, "history through Jan", f"{len(hist)} rows")
    max_observed = hist.observed_at.max()
    check("防泄漏-max(observed_at)<=decision", max_observed <= decision, f"<= {decision}", max_observed)
    check("防泄漏-2月1日未来真值未进入", not any(hist.date == date(2025, 2, 1)), False, any(hist.date == date(2025, 2, 1)))
    forecast = data_io.forecast_vintage_as_of(decision)
    latest_issue = forecast.issue_time.max()
    check("防泄漏-0点不能使用当日6点预报", latest_issue <= decision, f"<= {decision}", latest_issue)
    replay = data_io.replay_actual_day(date(2025, 2, 1))
    check("防泄漏-执行回放可读完整当日", len(replay) == 144, 144, len(replay))


def test_executor_and_auditor() -> None:
    slot = build_day_slots(date(2025, 1, 1))[0]
    curtail = executor.execute_q1(
        grid_contract=np.array([0.0]), price=np.ones(1), load=np.array([20.0]),
        pv_available=np.array([100.0]), pv_used=np.array([20.0]),
        pv_curtail=np.array([80.0]), charge=np.zeros(1), discharge=np.zeros(1),
        soc0=5000.0, starts=[slot.interval_start], ends=[slot.interval_end],
        plan_issue_time=datetime(2025, 1, 1),
    )
    check("执行器-弃光值原样传递", curtail[0].pv_used == 20 and curtail[0].pv_curtail == 80,
          "used=20,curtail=80", f"used={curtail[0].pv_used},curtail={curtail[0].pv_curtail}")
    check("核算器-合法弃光通过", accountant.audit(curtail).ok, True, accountant.audit(curtail).violations)
    expect_value_error("核算器-空记录拒绝", lambda: accountant.audit([]))

    base = make_records()
    broken = copy.deepcopy(base)
    broken[1] = replace(broken[1], soc_start=broken[1].soc_start + 1, soc_end=broken[1].soc_end + 1)
    check("核算器-SOC断链拒绝", not accountant.audit(broken).ok, False, accountant.audit(broken).violations)
    duplicate = [base[0], replace(base[1], slot=0), base[2]]
    check("核算器-重复槽拒绝", not accountant.audit(duplicate).ok, False, accountant.audit(duplicate).violations)
    gap = [base[0], replace(base[1], slot=2), replace(base[2], slot=3)]
    check("核算器-跳号拒绝", not accountant.audit(gap).ok, False, accountant.audit(gap).violations)
    reverse = list(reversed(base))
    check("核算器-逆序拒绝", not accountant.audit(reverse).ok, False, accountant.audit(reverse).violations)
    simultaneous = copy.deepcopy(base)
    simultaneous[0] = replace(
        simultaneous[0], grid_contract=11.9, grid_delivered=11.9,
        charge=10.0, discharge=8.1, normal_cost=11.9,
    )
    simultaneous_audit = accountant.audit(simultaneous)
    check("核算器-同时充放电拒绝", simultaneous_audit.simultaneous_charge_discharge == 1 and not simultaneous_audit.ok,
          "simultaneous=1 and rejected", f"simultaneous={simultaneous_audit.simultaneous_charge_discharge}; {simultaneous_audit.violations}")

    price = np.array([0.5, 1.0, 1.5]); load = np.array([400.0, 500.0, 300.0]); pv = np.array([0.0, 200.0, 0.0])
    sol = lp_kernel.solve_lp(lp_kernel.LpInputs(price, load, pv, 6000, 6000))
    slots = build_day_slots(date(2025, 1, 1))[:3]
    recs = executor.execute_q1(sol.grid, price, load, pv, sol.pv_used, sol.pv_curtail,
                               sol.charge, sol.discharge, 6000,
                               [s.interval_start for s in slots], [s.interval_end for s in slots], datetime(2025, 1, 1))
    aud = accountant.audit(recs)
    check("核算器-目标费用独立一致", abs(sol.cost - aud.cost_total) < 1e-6, sol.cost, aud.cost_total)
    baseline = lp_kernel.no_storage_schedule(price, load, pv, 6000)
    brecs = executor.execute_q1(baseline["grid"], price, load, pv, baseline["pv_used"], baseline["pv_curtail"],
                                baseline["charge"], baseline["discharge"], 6000,
                                [s.interval_start for s in slots], [s.interval_end for s in slots], datetime(2025, 1, 1))
    baud = accountant.audit(brecs)
    check("基线-无储能同终端SOC可行", baud.ok and abs(brecs[-1].soc_end - 6000) < 1e-9, True, baud.violations)
    check("基线-费用不低于LP", baud.cost_total + 1e-9 >= sol.cost, f">={sol.cost}", baud.cost_total)


def test_t4_milp() -> None:
    m = 5000 / 6
    constraints = [
        LinearConstraint([[1, 1, -1, 1, 0]], [50], [50]),
        LinearConstraint([[0, 0, 0.9, -1 / 0.9, 0]], [-np.inf], [0]),
        LinearConstraint([[0, 0, 1, 0, -m]], [-np.inf], [0]),
        LinearConstraint([[0, 0, 0, 1, m]], [-np.inf], [m]),
    ]
    result = milp(np.zeros(5), integrality=[0, 0, 0, 0, 1],
                  bounds=Bounds([100, 0, 0, 0, 0], [100, 0, m, m, 1]), constraints=constraints)
    check("T4-独立互斥MILP物理不可行", not result.success, "infeasible", result.message)


def workbook_snapshot(path: Path) -> tuple[list[tuple], dict[tuple, object]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    meta = [(ws.title, ws.sheet_state, ws.max_row, ws.max_column) for ws in wb.worksheets]
    cells = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title == "计划购电量" and cell.column == 2 and 2 <= cell.row <= 145:
                    continue
                cells[(ws.title, cell.coordinate)] = cell.value
    wb.close()
    return meta, cells


def test_official_template() -> None:
    source = Path("data/raw/official/附件5/result1.xlsm")
    before = workbook_snapshot(source)
    values = [k + 0.125 for k in range(144)]
    with tempfile.TemporaryDirectory(prefix="cumcm_template_") as tmp:
        output = Path(tmp) / "result1.xlsm"
        template_io.write_result1_plan(source, output, date(2025, 1, 1), values)
        back = template_io.read_result1_plan(output, date(2025, 1, 1))
        check("模板-官方副本144槽回读", np.allclose(back, values), values[:3], back[:3])
        after = workbook_snapshot(output)
        check("模板-工作表名称顺序隐藏状态保持", before[0] == after[0], before[0], after[0])
        check("模板-非目标单元格保持", before[1] == after[1], "unchanged", "unchanged" if before[1] == after[1] else "changed")
        wb = openpyxl.load_workbook(output, keep_vba=True)
        has_vba = wb.vba_archive is not None
        wb.close()
        check("模板-VBA容器保留", has_vba, True, has_vba)
        corrupt = Path(tmp) / "bad_label.xlsm"
        shutil.copy2(output, corrupt)
        wb = openpyxl.load_workbook(corrupt, keep_vba=True)
        wb["计划购电量"].cell(2, 1, "0:20-0:30")
        wb.save(corrupt); wb.close()
        expect_value_error("模板-错误标签拒绝", lambda: template_io.read_result1_plan(corrupt, date(2025, 1, 1)))
        expect_value_error("模板-少一槽拒绝", lambda: template_io.write_result1_plan(source, output, date(2025, 1, 1), values[:-1]))
        expect_value_error("模板-多一槽拒绝", lambda: template_io.write_result1_plan(source, output, date(2025, 1, 1), values + [1]))


def test_q1_anchor() -> None:
    frame = data_io.build_q1_inputs(date(2025, 1, 1), datetime(2025, 1, 1))
    inp = lp_kernel.LpInputs(frame.price.to_numpy(), frame.load_forecast.to_numpy() / 6,
                             frame.pv_forecast.to_numpy() / 6, 6000, 6000)
    ds = lp_kernel.solve_lp(inp, "highs-ds")
    ipm = lp_kernel.solve_lp(inp, "highs-ipm")
    check("Q1-双引擎一致", abs(ds.cost - ipm.cost) < 1e-6, ds.cost, ipm.cost)
    check("Q1-历史锚点保持", abs(ds.cost - 35126.948589) < 1e-4, 35126.948589, ds.cost)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    test_efficiency_and_lexicographic()
    test_dates_and_causality()
    test_executor_and_auditor()
    test_t4_milp()
    test_official_template()
    test_q1_anchor()
    failures = [r for r in RESULTS if not r["passed"]]
    payload = {
        "environment": {
            "python": __import__("sys").version.split()[0],
            "numpy": np.__version__, "scipy": scipy.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "summary": {"total": len(RESULTS), "passed": len(RESULTS) - len(failures), "failed": len(failures)},
        "results": RESULTS,
    }
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
