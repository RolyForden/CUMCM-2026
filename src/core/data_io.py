"""官方附件读取与因果数据访问接口。

Q1 输入、历史实际值、执行回放真值和预报批次使用不同入口。预测器只能
调用 history_actuals_as_of / forecast_vintage_as_of，不能把附件已经下载
等同于模拟决策时已经观测到未来真值。
"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from numbers import Real
from pathlib import Path

import openpyxl
import pandas as pd

from core.slot_adapter import N_SLOTS, build_day_slots

RAW = Path("data/raw")


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = pd.to_datetime(str(value), errors="raise")
    return parsed.date()


def _time_labels_xlsx(path: str | Path, sheet: str) -> list[str]:
    """读取附件1纵表第2至145行的时间标签。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    raw_labels = [
        row[0]
        for row in ws.iter_rows(min_row=2, max_row=145, min_col=1, max_col=1, values_only=True)
    ]
    wb.close()
    labels = []
    for value in raw_labels:
        if isinstance(value, time):
            labels.append(value.strftime("%H:%M"))
        else:
            labels.append(str(value).strip())
    if len(labels) != N_SLOTS:
        raise ValueError(f"{path}[{sheet}] 时间标签数不是144")
    return labels


def _read_wide(path: str | Path, sheet: str, day: date) -> pd.DataFrame:
    """按日期列精确读取某日144点；缺失或重复日期均报错。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    matches: list[tuple] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        if _as_date(row[0]) == day:
            matches.append(row)
    wb.close()
    if len(matches) != 1:
        raise ValueError(f"{path}[{sheet}] 日期 {day} 匹配行数={len(matches)}，要求唯一")
    values = matches[0][1:145]
    extra_values = matches[0][145:]
    if (
        len(values) != N_SLOTS
        or any(not isinstance(v, Real) or isinstance(v, bool) for v in values)
        or any(v is not None for v in extra_values)
    ):
        raise ValueError(f"{path}[{sheet}] 日期 {day} 必须恰含144个数值")

    slots = build_day_slots(day)
    return pd.DataFrame(
        {
            "date": [day] * N_SLOTS,
            "slot": [s.slot_id for s in slots],
            "interval_start": [s.interval_start for s in slots],
            "interval_end": [s.interval_end for s in slots],
            "valid_time": [s.interval_start for s in slots],
            "observed_at": [s.interval_end for s in slots],
            "value": [float(v) for v in values],
        }
    )


def attachment2_load(day: date) -> pd.DataFrame:
    return _read_wide(RAW / "official/附件2.xlsm", "小区负载", day)


def attachment2_pv(day: date) -> pd.DataFrame:
    return _read_wide(RAW / "substitute/附件2.xlsx", "光伏发电实际功率", day)


def attachment4(day: date) -> pd.DataFrame:
    return _read_wide(RAW / "official/附件4.xlsm", "Sheet1 (3)", day)


def build_q1_inputs(day: date, decision_time: datetime) -> pd.DataFrame:
    """读取附件1的已知单日价格、负荷和光伏预测。"""
    wb = openpyxl.load_workbook(
        RAW / "official/附件1.xlsm", read_only=True, data_only=True
    )
    ws = wb["Sheet1 (2)"]
    rows = list(ws.iter_rows(min_row=2, max_row=145, values_only=True))
    wb.close()
    if len(rows) != N_SLOTS:
        raise ValueError("附件1数据行数不是144")

    slots = build_day_slots(day)
    out = pd.DataFrame(
        {
            "date": [day] * N_SLOTS,
            "slot": [s.slot_id for s in slots],
            "interval_start": [s.interval_start for s in slots],
            "interval_end": [s.interval_end for s in slots],
            "valid_time": [s.interval_start for s in slots],
            "decision_time": [decision_time] * N_SLOTS,
            "price": [float(r[1]) for r in rows],
            "load_forecast": [float(r[2]) for r in rows],
            "pv_forecast": [float(r[3]) for r in rows],
        }
    )
    return out


def replay_actual_day(day: date) -> pd.DataFrame:
    """执行器/评价器专用：返回某日完整真值，不得传给预测器。"""
    load = attachment2_load(day).rename(columns={"value": "load_actual"})
    pv = attachment2_pv(day).rename(columns={"value": "pv_actual"})
    price = attachment4(day).rename(columns={"value": "price_actual"})
    keys = ["date", "slot", "interval_start", "interval_end", "valid_time", "observed_at"]
    out = load[keys + ["load_actual"]]
    out = out.merge(pv[["date", "slot", "pv_actual"]], on=["date", "slot"], validate="one_to_one")
    out = out.merge(price[["date", "slot", "price_actual"]], on=["date", "slot"], validate="one_to_one")
    return out


def history_actuals_as_of(decision_time: datetime) -> pd.DataFrame:
    """预测器专用：只返回在决策时刻前已经结束的实际槽位。"""
    start = date(2025, 1, 1)
    last_day = min(decision_time.date(), date(2025, 12, 31))
    if last_day < start:
        return pd.DataFrame()
    frames = []
    day = start
    while day <= last_day:
        frame = replay_actual_day(day)
        frame = frame[frame["observed_at"] <= decision_time]
        if not frame.empty:
            frames.append(frame)
        day += timedelta(days=1)
    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True)
    if out["observed_at"].max() > decision_time:
        raise AssertionError("历史接口泄漏决策时刻之后的实际值")
    return out


def forecast_vintage_as_of(decision_time: datetime) -> pd.DataFrame:
    """返回截至决策时刻已经发布的原始小时级光伏预报。

    这里只构造 issue_time、lead_hour 和小时有效时刻，不把小时值展开到
    10分钟槽位；N3首小时映射经人类裁决后再实现展开器。
    """
    wb = openpyxl.load_workbook(
        RAW / "official/附件3.xlsm", read_only=True, data_only=True
    )
    ws = wb["Sheet1 (2)"]
    rows = []
    current_day: date | None = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            current_day = _as_date(row[0])
        if current_day is None:
            raise ValueError("附件3首条预报缺少日期")
        hour_text = str(row[1]).strip()
        hour = int(hour_text.split(":", 1)[0])
        issue_time = datetime.combine(current_day, time(hour=hour))
        if issue_time > decision_time:
            continue
        for lead_hour, value in enumerate(row[2:26], start=1):
            rows.append(
                {
                    "issue_time": issue_time,
                    "valid_time": issue_time + timedelta(hours=lead_hour),
                    "lead_hour": lead_hour,
                    "pv_forecast": float(value),
                    "vintage": issue_time.isoformat(),
                }
            )
    wb.close()
    out = pd.DataFrame(rows)
    if not out.empty and out["issue_time"].max() > decision_time:
        raise AssertionError("预报接口泄漏尚未发布的批次")
    return out
