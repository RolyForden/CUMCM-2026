"""官方 result1 模板的原位填写与回读。"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl

from core.slot_adapter import N_SLOTS, build_day_slots


def write_result1_plan(
    template_path: str | Path,
    output_path: str | Path,
    day: date,
    grid: list[float],
) -> None:
    """复制官方模板语义，仅填写“计划购电量”表的144个数值。"""
    if len(grid) != N_SLOTS:
        raise ValueError(f"购电计划长度 {len(grid)} != 144")
    keep_vba = str(template_path).lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(template_path, keep_vba=keep_vba)
    if "计划购电量" not in wb.sheetnames:
        wb.close()
        raise ValueError("官方模板缺少计划购电量工作表")
    ws = wb["计划购电量"]
    slots = build_day_slots(day)
    for k, slot in enumerate(slots):
        label = ws.cell(row=k + 2, column=1).value
        if str(label).strip() != slot.template_interval:
            wb.close()
            raise ValueError(
                f"slot {k}: 官方模板标签 {label!r} != {slot.template_interval!r}"
            )
        ws.cell(row=k + 2, column=2, value=float(grid[k]))
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()


def read_result1_plan(path: str | Path, day: date) -> list[float]:
    """从官方结构的“计划购电量”表回读144槽。"""
    keep_vba = str(path).lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=keep_vba)
    if "计划购电量" not in wb.sheetnames:
        wb.close()
        raise ValueError("结果文件缺少计划购电量工作表")
    ws = wb["计划购电量"]
    slots = build_day_slots(day)
    values = []
    for k, slot in enumerate(slots):
        label = ws.cell(row=k + 2, column=1).value
        value = ws.cell(row=k + 2, column=2).value
        if str(label).strip() != slot.template_interval:
            wb.close()
            raise ValueError(f"slot {k}: 结果标签错位")
        if value is None:
            wb.close()
            raise ValueError(f"slot {k}: 购电量为空")
        values.append(float(value))
    wb.close()
    return values
