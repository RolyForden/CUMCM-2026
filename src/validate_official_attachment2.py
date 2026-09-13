"""最小验证：读取官方附件2的负荷/光伏表，并与旧光伏副本逐值比较。"""

from __future__ import annotations

import argparse
from datetime import date
import json
from numbers import Real
from pathlib import Path
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from core.data_io import _read_wide


PV_SHEET = "光伏发电实际功率"
LOAD_SHEET = "小区负载"


def read_matrix(path: Path, sheet: str) -> list[float]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{path} 缺少工作表 {sheet}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    if len(rows) != 365:
        raise ValueError(f"{path}[{sheet}] 日数={len(rows)}，要求365")
    values = [value for row in rows for value in row[1:145]]
    if len(values) != 365 * 144 or any(
        not isinstance(value, Real) or isinstance(value, bool) for value in values
    ):
        raise ValueError(f"{path}[{sheet}] 必须恰含52560个数值")
    return [float(value) for value in values]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--official",
        type=Path,
        default=ROOT / "data/raw/official/附件2.xlsx",
    )
    parser.add_argument(
        "--legacy",
        type=Path,
        help="可选：旧光伏核对副本；提供时执行52560值全量比较",
    )
    args = parser.parse_args()

    # 通过生产宽表读取器检查首末日，防止只比较原始单元格。
    for day in (date(2025, 1, 1), date(2025, 12, 31)):
        if len(_read_wide(args.official, LOAD_SHEET, day)) != 144:
            raise AssertionError(f"{day} 负荷读取数不是144")
        if len(_read_wide(args.official, PV_SHEET, day)) != 144:
            raise AssertionError(f"{day} 光伏读取数不是144")

    official = read_matrix(args.official, PV_SHEET)
    result = {
        "official": str(args.official),
        "pv_value_count": len(official),
    }
    if args.legacy is not None:
        legacy = read_matrix(args.legacy, PV_SHEET)
        max_abs_diff = max(abs(a - b) for a, b in zip(official, legacy, strict=True))
        result.update(
            legacy=str(args.legacy),
            max_abs_diff=max_abs_diff,
            identical=max_abs_diff == 0.0,
        )
        if not result["identical"]:
            raise AssertionError(json.dumps(result, ensure_ascii=False))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
