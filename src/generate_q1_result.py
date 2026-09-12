"""按D002与D003方案A生成并审计Q1正式结果。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
import sys
from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl

from core import accountant, data_io, executor, lp_kernel, template_io
from core.slot_adapter import build_day_slots


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = PROJECT_ROOT / "data/raw/official/附件5/result1.xlsm"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "outputs/q1_scheme_a"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def workbook_snapshot(path: Path, official_titles: tuple[str, ...]) -> dict[str, Any]:
    """记录结构与非目标单元格；B2:B145是唯一允许变化的范围。

    official_titles：官方模板已有工作表，只对这些表做非目标单元格比对；
    新增的“充放电量”表不参与该比对（其内容由回读校验覆盖）。
    """
    wb = openpyxl.load_workbook(path, keep_vba=True, data_only=False)
    sheets = []
    non_target_cells = []
    for ws in wb.worksheets:
        sheets.append(
            {
                "title": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
            }
        )
        if ws.title not in official_titles:
            continue
        for row in ws.iter_rows():
            for cell in row:
                is_target = ws.title == "计划购电量" and cell.column == 2 and 2 <= cell.row <= 145
                if not is_target:
                    non_target_cells.append(
                        [ws.title, cell.coordinate, cell.value, cell.data_type, cell.number_format]
                    )
    # keep_vba 的 archive 包含工作簿全部 ZIP 条目；这里只比较真正与宏有关的
    # 部件。普通 sharedStrings/目录项会被 openpyxl 合法重写，不能冒充 VBA。
    archive_names = wb.vba_archive.namelist() if wb.vba_archive is not None else []
    vba_names = sorted(
        name for name in archive_names
        if "vba" in name.lower() or "macro" in name.lower()
    )
    wb.close()
    return {"sheets": sheets, "non_target_cells": non_target_cells, "vba_names": vba_names}


def git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=PROJECT_ROOT, text=True
    ).strip()


def write_dispatch_csv(path: Path, recs: list[executor.ExecRecord]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(recs[0]).keys())
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for record in recs:
            row = asdict(record)
            row["interval_start"] = record.interval_start.isoformat()
            row["interval_end"] = record.interval_end.isoformat()
            writer.writerow(row)


def build_result(template: Path, output_dir: Path) -> dict[str, Any]:
    day = date(2025, 1, 1)
    decision_time = datetime(2025, 1, 1, 0, 0)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / "result1.xlsm"
    dispatch_path = output_dir / "q1_dispatch.csv"
    paper_tables_path = output_dir / "q1_paper_tables.json"
    audit_path = output_dir / "q1_audit.json"

    official_titles = tuple(
        openpyxl.load_workbook(template, read_only=True).sheetnames
    )
    source_hash_before = sha256(template)
    source_snapshot = workbook_snapshot(template, official_titles)
    frame = data_io.build_q1_inputs(day, decision_time)
    inputs = lp_kernel.LpInputs(
        price=frame["price"].to_numpy(dtype=float),
        load=frame["load_forecast"].to_numpy(dtype=float) * lp_kernel.DT,
        pv_available=frame["pv_forecast"].to_numpy(dtype=float) * lp_kernel.DT,
        soc0=6000.0,
        soc_final=6000.0,
    )
    ds = lp_kernel.solve_lp(inputs, solver="highs-ds")
    ipm = lp_kernel.solve_lp(inputs, solver="highs-ipm")
    if ds.status != 0 or ipm.status != 0:
        raise RuntimeError(f"Q1求解失败：DS={ds.message}; IPM={ipm.message}")

    slots = build_day_slots(day)
    recs = executor.execute_q1(
        grid_contract=ds.grid,
        price=inputs.price,
        load=inputs.load,
        pv_available=inputs.pv_available,
        pv_used=ds.pv_used,
        pv_curtail=ds.pv_curtail,
        charge=ds.charge,
        discharge=ds.discharge,
        soc0=6000.0,
        starts=[slot.interval_start for slot in slots],
        ends=[slot.interval_end for slot in slots],
        plan_issue_time=decision_time,
    )
    audit = accountant.audit(recs)

    # 6 个 4 小时块汇总（题面表 2 布局；按槽位顺序每 24 槽一块）
    block_charge = [float(np.sum(ds.charge[i * 24:(i + 1) * 24])) for i in range(6)]
    block_discharge = [float(np.sum(ds.discharge[i * 24:(i + 1) * 24])) for i in range(6)]

    template_io.write_result1_plan(template, workbook_path, day, ds.grid.tolist())
    template_io.add_result1_cd_sheet(
        workbook_path,
        block_charge,
        block_discharge,
        soc_0000=recs[0].soc_start,
        soc_2400=recs[142].soc_end,
    )
    readback = np.asarray(template_io.read_result1_plan(workbook_path, day))
    cd_readback = template_io.read_result1_cd_sheet(workbook_path)
    output_snapshot = workbook_snapshot(workbook_path, official_titles)
    source_hash_after = sha256(template)
    readback_max_abs_diff = float(np.max(np.abs(readback - ds.grid)))
    cd_block_max_abs_diff = max(
        max(abs(a - b) for a, b in zip(cd_readback["blocks"][i],
                                       (block_charge[i], block_discharge[i])))
        for i in range(6)
    )
    soc_e143 = recs[142].soc_end

    # 论文表 1：指定时间段购电量 + 全天购电量/购电费
    # 方案 A：区间 [hh:mm, hh:mm+10) 的槽位 = 分钟数/10 − 1
    # 10:00 → 60 分钟*10 → slot 59；12:00 → 71；…；20:00 → 119
    paper_table1_slots = {
        "10:00-10:10": 59,
        "12:00-12:10": 71,
        "14:00-14:10": 83,
        "16:00-16:10": 95,
        "18:00-18:10": 107,
        "20:00-20:10": 119,
    }
    table1 = {
        label: float(ds.grid[k])
        for label, k in paper_table1_slots.items()
    }
    table1["全天购电量"] = float(np.sum(ds.grid))
    table1["全天购电费"] = audit.cost_total

    paper_tables = {
        "table1": table1,
        "table2": {
            "blocks": [
                {"label": template_io.CD_BLOCK_LABELS[i],
                 "covers_slots": f"{i*24}-{i*24+23}",
                 "charge": block_charge[i],
                 "discharge": block_discharge[i]}
                for i in range(6)
            ],
            "soc_0000": recs[0].soc_start,
            "soc_2400": soc_e143,
        },
    }

    checks = {
        "dual_simplex_feasible": ds.status == 0,
        "ipm_feasible": ipm.status == 0,
        "dual_engine_cost_match": abs(ds.cost - ipm.cost) < 1e-6,
        "independent_audit_ok": audit.ok,
        "optimizer_accountant_cost_match": abs(ds.cost - audit.cost_total) < 1e-6,
        "max_energy_residual_below_1e-9": audit.max_residual < 1e-9,
        "terminal_soc_e144_is_6000": abs(recs[-1].soc_end - 6000.0) < 1e-6,
        "no_simultaneous_charge_discharge": audit.simultaneous_charge_discharge == 0,
        "template_readback_144_slots": len(readback) == 144,
        "template_readback_matches_plan": readback_max_abs_diff < 1e-9,
        "cd_sheet_readback_matches_blocks": cd_block_max_abs_diff < 1e-9,
        "cd_sheet_soc_matches": (
            abs(cd_readback["soc_0000"] - recs[0].soc_start) < 1e-9
            and abs(cd_readback["soc_2400"] - soc_e143) < 1e-9
        ),
        "official_sheets_untouched": all(
            s in output_snapshot["sheets"]
            for s in source_snapshot["sheets"]
        ),
        "template_non_target_cells_preserved": source_snapshot["non_target_cells"] == output_snapshot["non_target_cells"],
        "template_vba_container_preserved": source_snapshot["vba_names"] == output_snapshot["vba_names"],
        "raw_template_unchanged": source_hash_before == source_hash_after,
    }
    checks = {name: bool(passed) for name, passed in checks.items()}
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise RuntimeError(f"Q1正式结果验收失败：{failed}")

    write_dispatch_csv(dispatch_path, recs)
    paper_tables_path.write_text(
        json.dumps(paper_tables, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    evidence = {
        "status": "PASS",
        "scheme": "D003方案A：slot 0从00:10开始，完整执行144槽，E_144=E_0=6000 kWh",
        "model_commit": git_head(),
        "generated_at": datetime.now().astimezone().isoformat(),
        "inputs": {
            "day": day.isoformat(),
            "decision_time": decision_time.isoformat(),
            "template": str(template.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "template_sha256": source_hash_before,
        },
        "outputs": {
            "workbook": str(workbook_path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "workbook_sha256": sha256(workbook_path),
            "dispatch_csv": str(dispatch_path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            "dispatch_csv_sha256": sha256(dispatch_path),
            "paper_tables": str(paper_tables_path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
        },
        "solver": {
            "highs_ds_cost": ds.cost,
            "highs_ds_primary_optimum": ds.primary_optimum,
            "highs_ipm_cost": ipm.cost,
            "cost_difference": abs(ds.cost - ipm.cost),
        },
        "audit": asdict(audit) | {"ok": audit.ok},
        "summary": {
            "grid_total_kwh": float(np.sum(ds.grid)),
            "pv_used_total_kwh": float(np.sum(ds.pv_used)),
            "pv_curtail_total_kwh": float(np.sum(ds.pv_curtail)),
            "charge_total_kwh": float(np.sum(ds.charge)),
            "discharge_total_kwh": float(np.sum(ds.discharge)),
            "soc_e0_kwh": recs[0].soc_start,
            "soc_e143_kwh": soc_e143,
            "soc_e144_kwh": recs[143].soc_end,
            "readback_max_abs_diff": readback_max_abs_diff,
            "cd_readback_max_abs_diff": cd_block_max_abs_diff,
        },
        "paper_tables": paper_tables,
        "template_snapshot": {
            "official_source_sheets": source_snapshot["sheets"],
            "output_sheets": output_snapshot["sheets"],
        },
        "checks": checks,
    }
    audit_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    return evidence


def reopen_verify(output_dir: Path, evidence: dict[str, Any]) -> dict[str, bool]:
    """独立重开核验：只读 evidence + 产物文件，不复用生成期对象。

    1. 从 q1_dispatch.csv 独立复算能量残差、SOC 链、费用（与审计同口径，
       但不引用优化器/执行器对象）；
    2. 重新打开 result1.xlsm 回读计划购电量表与充放电量表；
    3. 对照 evidence 中的 SHA-256。
    """
    dispatch_path = output_dir / "q1_dispatch.csv"
    workbook_path = output_dir / "result1.xlsm"
    checks: dict[str, bool] = {}
    with dispatch_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != 144:
        raise RuntimeError(f"dispatch.csv 行数 {len(rows)} != 144")
    e = 0.0
    max_res = 0.0
    total_cost = 0.0
    for i, r in enumerate(rows):
        grid = float(r["grid_delivered"])
        pv_used = float(r["pv_used"])
        pv_curtail = float(r["pv_curtail"])
        pv_avail = float(r["pv_available"])
        load = float(r["load"])
        charge = float(r["charge"])
        discharge = float(r["discharge"])
        price = float(r["price"])
        soc_start = float(r["soc_start"])
        soc_end = float(r["soc_end"])
        e_next = soc_start + 0.9 * charge - discharge / 0.9
        if i > 0:
            prev_end = float(rows[i - 1]["soc_end"])
            if abs(prev_end - soc_start) > 1e-6:
                raise RuntimeError(f"重开核验：slot {i} SOC 断链")
        max_res = max(max_res, abs(grid + pv_used + discharge - load - charge))
        if abs(e_next - soc_end) > 1e-6:
            raise RuntimeError(f"重开核验：slot {i} SOC 转移不符")
        if pv_used + pv_curtail - pv_avail > 1e-6 or pv_used > pv_avail + 1e-9:
            raise RuntimeError(f"重开核验：slot {i} 光伏守恒失败")
        if charge > 1e-6 and discharge > 1e-6:
            raise RuntimeError(f"重开核验：slot {i} 同时充放电")
        total_cost += price * grid
    checks["csv_144_slots"] = True
    checks["csv_max_residual_below_1e-9"] = max_res < 1e-9
    checks["csv_total_cost_matches_audit"] = (
        abs(total_cost - evidence["audit"]["cost_total"]) < 1e-6
    )
    checks["csv_sha256_matches_evidence"] = (
        sha256(dispatch_path) == evidence["outputs"]["dispatch_csv_sha256"]
    )
    checks["workbook_sha256_matches_evidence"] = (
        sha256(workbook_path) == evidence["outputs"]["workbook_sha256"]
    )
    day = date.fromisoformat(evidence["inputs"]["day"])
    plan = template_io.read_result1_plan(workbook_path, day)
    grid_csv = [float(r["grid_delivered"]) for r in rows]
    checks["reopen_plan_sheet_matches_csv"] = (
        max(abs(a - b) for a, b in zip(plan, grid_csv)) < 1e-9
    )
    cd = template_io.read_result1_cd_sheet(workbook_path)
    blocks = evidence["paper_tables"]["table2"]["blocks"]
    checks["reopen_cd_sheet_matches_evidence"] = all(
        abs(cd["blocks"][i][0] - blocks[i]["charge"]) < 1e-9
        and abs(cd["blocks"][i][1] - blocks[i]["discharge"]) < 1e-9
        for i in range(6)
    )
    checks["reopen_cd_soc_matches_evidence"] = (
        abs(cd["soc_0000"] - evidence["paper_tables"]["table2"]["soc_0000"]) < 1e-9
        and abs(cd["soc_2400"] - evidence["paper_tables"]["table2"]["soc_2400"]) < 1e-9
    )
    checks = {k: bool(v) for k, v in checks.items()}
    if not all(checks.values()):
        raise RuntimeError(f"重开核验失败：{[k for k, v in checks.items() if not v]}")
    return checks


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    evidence = build_result(args.template.resolve(), output_dir)
    reopen_checks = reopen_verify(output_dir, evidence)
    audit_path = output_dir / "q1_audit.json"
    stored = json.loads(audit_path.read_text(encoding="utf-8"))
    stored["reopen_checks"] = reopen_checks
    audit_path.write_text(json.dumps(stored, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "status": evidence["status"],
        "cost": evidence["audit"]["cost_total"],
        "checks_passed": sum(evidence["checks"].values()),
        "checks_total": len(evidence["checks"]),
        "reopen_checks_passed": sum(reopen_checks.values()),
        "reopen_checks_total": len(reopen_checks),
        "workbook": evidence["outputs"]["workbook"],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
